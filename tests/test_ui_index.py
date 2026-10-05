"""careeros.ui.index: the SQLite read index (data/careeros.db) rebuilt from the job, run and tracker files."""
from __future__ import annotations

import json
import os
import sqlite3
from datetime import datetime, timezone

import pytest
from conftest import make_temp_root
from fixtures.ui_data import build_ui_data

from careeros.store import Store
from careeros.ui import index as index_mod
from careeros.ui.index import Index

pytestmark = pytest.mark.unit

NOW = datetime(2026, 9, 24, 15, 0, tzinfo=timezone.utc)  # a Thursday


@pytest.fixture
def data(tmp_path):
    return build_ui_data(make_temp_root(tmp_path / "repo"), NOW)


@pytest.fixture
def idx(data):
    ix = Index(data["settings"])
    ix.rebuild()
    yield ix
    ix.close()


def _jobs(ix):
    return {r["job_id"]: r for r in ix.query("SELECT * FROM jobs")}


def test_default_path_is_next_to_data_jobs(data):
    ix = Index(data["settings"])
    assert ix.path == data["settings"].paths["jobs_dir"].parent / "careeros.db"
    ix.close()


def test_rebuild_indexes_every_job_with_its_fields(idx, data):
    jobs = _jobs(idx)
    assert set(jobs) == set(data["jobs"].values())
    r = jobs[data["jobs"]["review"]]
    assert (r["company"], r["title"], r["status"], r["fit"], r["tier"], r["safety"], r["category"]) == (
        "Umbrella Labs", "Infrastructure Engineer", "needs_review", 91, "A", "review", "swe_backend")
    assert r["qa_passed"] == 1 and r["qa_score"] == pytest.approx(8.2)
    assert jobs[data["jobs"]["found"]]["fit"] is None and jobs[data["jobs"]["found"]]["qa_passed"] is None
    assert jobs[data["jobs"]["applied"]]["applied_at"].startswith("2026-09-22")
    assert jobs[data["jobs"]["found"]]["applied_at"] is None


def test_status_history_rows(idx, data):
    rows = idx.query("SELECT status FROM status_history WHERE job_id = ? ORDER BY seq", (data["jobs"]["interview"],))
    assert [r["status"] for r in rows] == ["found", "scored", "queued", "applied", "screening", "interview"]


def test_action_items_and_contacts_and_runs(idx, data):
    acts = {r["id"]: r for r in idx.query("SELECT * FROM action_items")}
    assert set(acts) == set(data["actions"].values())
    assert acts[data["actions"]["done"]]["done"] == 1 and acts[data["actions"]["high"]]["priority"] == "H"
    assert acts[data["actions"]["high"]]["link"] == "https://boards.example.com/review"
    contacts = idx.query("SELECT * FROM contacts ORDER BY name")
    assert [(c["name"], c["linkedin_degree"], c["mutuals"]) for c in contacts] == [("Pat Rivers", 1, None),
                                                                                   ("Sam Lee", None, 0)]
    runs = {r["id"]: r for r in idx.query("SELECT * FROM runs")}
    assert runs[data["runs"]["prepare"]]["stop_reason"] == "usage_limit"
    att = idx.query("SELECT * FROM attempts WHERE run_id = ?", (data["runs"]["score"],))
    assert att[0]["outcome"] == "ok" and att[0]["job_id"] == data["jobs"]["queued"]


def test_action_item_due_is_indexed(idx, data):
    acts = {r["id"]: r for r in idx.query("SELECT * FROM action_items")}
    high = acts[data["actions"]["high"]]
    assert high["due"] and high["due"].startswith("20") and high["due_reason"] == "posting closes"
    assert acts[data["actions"]["medium"]]["due"] is None


def test_wal_mode_and_schema_version(idx):
    assert idx.query("PRAGMA journal_mode")[0]["journal_mode"] == "wal"
    assert idx.get_meta("schema_version") == str(index_mod.SCHEMA_VERSION)
    assert idx.get_meta("indexed_at")


def test_sync_skips_unchanged_jobs_and_picks_up_changes(idx, data):
    assert idx.sync()["jobs_changed"] == []
    jid = data["jobs"]["queued"]
    Store(data["settings"]).set_status(jid, "prepared", "docs ready")
    assert idx.sync()["jobs_changed"] == [jid]
    assert _jobs(idx)[jid]["status"] == "prepared"


def test_update_jobs_by_id(idx, data):
    jid = data["jobs"]["found"]
    Store(data["settings"]).set_status(jid, "scored")
    assert idx.update_jobs([jid]) == [jid]
    assert idx.update_jobs([jid]) == []          # same files -> skipped
    assert _jobs(idx)[jid]["status"] == "scored"


def test_deleted_job_leaves_the_index(idx, data):
    import shutil

    jid = data["jobs"]["skipped"]
    shutil.rmtree(Store(data["settings"]).job_dir(jid))
    res = idx.sync()
    assert jid in res["jobs_removed"] and jid not in _jobs(idx)
    assert idx.query("SELECT COUNT(*) AS n FROM status_history WHERE job_id = ?", (jid,))[0]["n"] == 0
    other = data["jobs"]["rejected"]                                     # by id (the watcher's path): also removes
    shutil.rmtree(Store(data["settings"]).job_dir(other))
    assert idx.update_jobs([other]) == [other] and other not in _jobs(idx)
    assert idx.update_jobs([other]) == []


def test_finder_copies_are_ignored(idx, data):
    import shutil

    jd = Store(data["settings"]).job_dir(data["jobs"]["found"])
    shutil.copytree(jd, jd.with_name(jd.name + " 2"))
    idx.sync()
    assert not any(" " in j for j in _jobs(idx))


def test_schema_bump_rebuilds(data, monkeypatch):
    ix = Index(data["settings"])
    ix.rebuild()
    ix.close()
    con = sqlite3.connect(Index(data["settings"]).path)
    con.execute("UPDATE meta SET value = '0' WHERE key = 'schema_version'")
    con.execute("DELETE FROM jobs")
    con.commit()
    con.close()
    ix = Index(data["settings"])
    ix.sync()
    assert len(_jobs(ix)) == len(data["jobs"]) and ix.get_meta("schema_version") == str(index_mod.SCHEMA_VERSION)
    ix.close()


def test_tracker_reindexed_only_when_it_changes(idx, data):
    from careeros.tracker import Tracker

    assert idx.update_tracker() is False
    Tracker(settings=data["settings"]).mark_action_done(data["actions"]["high"])
    st = data["settings"].paths["tracker_xlsx"]
    os.utime(st, (st.stat().st_atime, st.stat().st_mtime + 5))
    assert idx.update_tracker() is True
    assert idx.query("SELECT done FROM action_items WHERE id = ?", (data["actions"]["high"],))[0]["done"] == 1


def test_missing_tracker_is_empty_and_never_created(tmp_path):
    from careeros.config import Settings

    s = Settings.load(make_temp_root(tmp_path / "repo"))
    ix = Index(s)
    ix.rebuild()
    assert ix.query("SELECT COUNT(*) AS n FROM action_items")[0]["n"] == 0
    assert not s.paths["tracker_xlsx"].exists()
    assert ix.query("SELECT COUNT(*) AS n FROM jobs")[0]["n"] == 0
    ix.close()


def test_broken_json_does_not_stop_the_index(idx, data):
    jid = data["jobs"]["scored"]
    (Store(data["settings"]).job_dir(jid) / "score.json").write_text("{not json", encoding="utf-8")
    idx.update_jobs([jid])
    row = _jobs(idx)[jid]
    assert row["fit"] is None and row["company"] == "Globex"


def test_run_updates(idx, data):
    from careeros.runs.store import RunStore

    rs = RunStore(data["settings"])
    run = rs.load_run(data["runs"]["score"])
    run["stop_reason"] = "cancelled"
    rs.save_run(run)
    assert idx.update_runs([run["id"]]) == [run["id"]]
    assert idx.query("SELECT stop_reason FROM runs WHERE id = ?", (run["id"],))[0]["stop_reason"] == "cancelled"
    assert idx.sync()["runs_changed"] == []


def test_contacts_json_parse_errors_are_skipped(idx, data):
    jd = Store(data["settings"]).job_dir(data["jobs"]["interview"])
    (jd / "contacts.json").write_text(json.dumps({"contacts": "nope"}), encoding="utf-8")
    idx.update_jobs([data["jobs"]["interview"]])
    assert idx.query("SELECT COUNT(*) AS n FROM contacts")[0]["n"] == 0


@pytest.mark.parametrize("damage", ["garbage_pages", "truncated_index"])
def test_damaged_index_is_renamed_aside_and_rebuilt(data, damage):
    path = Index(data["settings"]).path
    if damage == "garbage_pages":
        Index.remove_files(path)
        path.write_bytes(b"SQLite format 3\x00" + b"damaged page" * 50)
    else:                                   # our own index, cut short mid-write
        ix = Index(data["settings"])
        ix.rebuild()
        ix.close()
        for side in ("-wal", "-shm"):
            path.with_name(path.name + side).unlink(missing_ok=True)
        path.write_bytes(path.read_bytes()[:1500])
    before = path.read_bytes()
    ix = Index(data["settings"])
    ix.sync()
    assert ix.query("SELECT COUNT(*) AS n FROM jobs")[0]["n"] == len(data["jobs"])
    ix.close()
    aside = list(path.parent.glob(path.name + ".corrupt-*"))
    assert len(aside) == 1 and aside[0].read_bytes() == before
    assert len(aside[0].name.rsplit(".corrupt-", 1)[1]) == len("20260924-150000")


def test_reindex_renames_a_damaged_index_aside(data):
    path = Index(data["settings"]).path
    Index.remove_files(path)
    path.write_bytes(b"SQLite format 3\x00" + b"x" * 400)
    before = path.read_bytes()
    Index.remove_files(path)
    assert not path.exists()
    aside = list(path.parent.glob(path.name + ".corrupt-*"))
    assert len(aside) == 1 and aside[0].read_bytes() == before


def test_empty_file_at_the_index_path_is_adopted(data):
    path = Index(data["settings"]).path
    Index.remove_files(path)
    path.write_bytes(b"")
    ix = Index(data["settings"])
    ix.sync()
    assert ix.query("SELECT COUNT(*) AS n FROM jobs")[0]["n"] == len(data["jobs"])
    ix.close()


def test_a_non_sqlite_file_at_the_index_path_is_never_deleted(data):
    from careeros.config import ConfigError

    path = Index(data["settings"]).path
    Index.remove_files(path)
    path.write_bytes(b"someone's notes, not a database")
    with pytest.raises(ConfigError):
        Index(data["settings"])
    with pytest.raises(ConfigError):
        Index.remove_files(path)
    assert path.read_bytes() == b"someone's notes, not a database"


def _with_index_path(data, value):
    s = data["settings"]
    s.pipeline.setdefault("ui", {})["index_path"] = value
    return s


@pytest.mark.parametrize("target", ["tracker", "folder", "config", "profile", "jobs_dir", "in_job", "in_runs"])
def test_index_path_may_not_point_at_user_files(data, target):
    from careeros.config import ConfigError
    from careeros.ui.index import default_path

    s = data["settings"]
    tracker = s.paths["tracker_xlsx"]
    before = tracker.read_bytes()
    value = {"tracker": str(tracker), "folder": str(s.root / "data"), "config": "config/ui.db",
             "profile": "profile/ui.db", "jobs_dir": str(s.paths["jobs_dir"]),
             "in_job": "data/jobs/abc123/score.json", "in_runs": "data/runs/x.db"}[target]
    _with_index_path(data, value)
    with pytest.raises(ConfigError):
        default_path(s)
    with pytest.raises(ConfigError):
        Index(s)
    assert tracker.read_bytes() == before


def test_runs_and_tracker_updates_bump_indexed_at(idx, data):
    from careeros.runs.store import RunStore

    idx.set_meta("indexed_at", "old")
    rs = RunStore(data["settings"])
    run = rs.load_run(data["runs"]["score"])
    run["detail"] = "edited"
    rs.save_run(run)
    assert idx.update_runs([run["id"]]) and idx.get_meta("indexed_at") != "old"
    idx.set_meta("indexed_at", "old")
    st = data["settings"].paths["tracker_xlsx"]
    os.utime(st, (st.stat().st_atime, st.stat().st_mtime + 9))
    assert idx.update_tracker() and idx.get_meta("indexed_at") != "old"
    idx.set_meta("indexed_at", "old")
    assert idx.update_runs([run["id"]]) == [] and idx.get_meta("indexed_at") == "old"


def test_deleted_run_folder_leaves_the_index(idx, data):
    import shutil

    from careeros.runs.store import RunStore

    rid = data["runs"]["prepare"]
    shutil.rmtree(RunStore(data["settings"]).run_dir(rid))
    res = idx.sync()
    assert res["runs_removed"] == [rid]
    assert idx.query("SELECT COUNT(*) AS n FROM runs WHERE id = ?", (rid,))[0]["n"] == 0
    assert idx.query("SELECT COUNT(*) AS n FROM attempts WHERE run_id = ?", (rid,))[0]["n"] == 0


@pytest.mark.parametrize("junk", [b"partial write", b"PK\x03\x04 truncated zip"])
def test_unreadable_tracker_keeps_old_rows_and_is_never_touched(idx, data, junk):
    tr = data["settings"].paths["tracker_xlsx"]
    before = idx.query("SELECT id FROM action_items ORDER BY id")
    tr.write_bytes(junk)
    assert idx.update_tracker() is False
    assert idx.query("SELECT id FROM action_items ORDER BY id") == before
    assert tr.read_bytes() == junk
    assert not list(tr.parent.glob("*corrupt*"))
    assert idx.update_tracker() is False           # signature not stored: retried, still failing


def test_tracker_is_read_without_the_tracker_class(idx, data, monkeypatch):
    import careeros.tracker as tracker_mod

    def boom(*a, **k):
        raise AssertionError("the index must not go through Tracker (it can rename or re-init the workbook)")

    monkeypatch.setattr(tracker_mod.Tracker, "_load", boom)
    st = data["settings"].paths["tracker_xlsx"]
    os.utime(st, (st.stat().st_atime, st.stat().st_mtime + 7))
    assert idx.update_tracker() is True
    assert len(idx.query("SELECT id FROM action_items")) == 4


def test_qa_review_schema_is_indexed(idx, data):
    jid = data["jobs"]["queued"]
    jd = Store(data["settings"]).job_dir(jid)
    (jd / "qa.json").write_text(json.dumps({"job_id": jid, "mean": 6.4, "pass": False, "rubric": {},
                                            "deterministic": {"pass": True}}), encoding="utf-8")
    idx.update_jobs([jid])
    row = _jobs(idx)[jid]
    assert row["qa_passed"] == 0 and row["qa_score"] == pytest.approx(6.4)


def test_legacy_qa_results_list_still_indexed(idx, data):
    from careeros.models import QAResult

    jid = data["jobs"]["scored"]
    Store(data["settings"]).save_qa(jid, [QAResult(job_id=jid, artifact="resume", passed=True,
                                                   critic_scores={"voice": 7.0})])
    idx.update_jobs([jid])
    row = _jobs(idx)[jid]
    assert row["qa_passed"] == 1 and row["qa_score"] == pytest.approx(7.0)


def test_excel_date_cell_due_is_indexed_as_a_date(idx, data):
    """Excel turns a typed date into a datetime at 00:00; that means the whole day, not midnight (overdue)."""
    from openpyxl import load_workbook

    from careeros.ui.services import actions as svc

    st = data["settings"].paths["tracker_xlsx"]
    wb = load_workbook(st)
    ws = wb["Action Items"]
    hdr = {c.value: c.column for c in ws[1] if c.value}
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, hdr["ID"]).value == data["actions"]["medium"])
    ws.cell(row, hdr["Due"]).value = datetime(2026, 9, 24)
    wb.save(st)
    os.utime(st, (st.stat().st_atime, st.stat().st_mtime + 5))
    assert idx.update_tracker() is True
    due = idx.query("SELECT due FROM action_items WHERE id = ?", (data["actions"]["medium"],))[0]["due"]
    assert due == "2026-09-24"
    assert svc.due_bucket(due, NOW, timezone.utc) == "today"


def test_dir_sig_changes_on_rename_and_same_size_edit(tmp_path):
    import os

    from careeros.ui.index import _dir_sig

    d = tmp_path / "job"
    d.mkdir()
    f = d / "a.json"
    f.write_text("aaaa", encoding="utf-8")
    os.utime(f, ns=(1_000_000_000, 1_000_000_000))
    s0 = _dir_sig(d)
    f.rename(d / "b.json")
    s1 = _dir_sig(d)
    assert s1 != s0
    (d / "b.json").write_text("bbbb", encoding="utf-8")
    os.utime(d / "b.json", ns=(1_000_000_000, 1_000_000_000))
    (d / "c.json").write_text("cc", encoding="utf-8")
    s2 = _dir_sig(d)
    (d / "c.json").write_text("dd", encoding="utf-8")
    os.utime(d / "c.json", ns=(2_000_000_000, 2_000_000_000))
    os.utime(d / "b.json", ns=(3_000_000_000, 3_000_000_000))
    s3 = _dir_sig(d)
    assert s3 != s2
    (d / "x.json.tmp").write_text("partial", encoding="utf-8")
    assert _dir_sig(d) == s3


def _foreign_sqlite(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE jobs (id INTEGER, title TEXT)")
    con.execute("INSERT INTO jobs VALUES (1, 'kept')")
    con.commit()
    con.close()
    return path.read_bytes()


def test_another_apps_sqlite_file_is_never_adopted(data):
    from careeros.config import ConfigError

    path = Index(data["settings"]).path
    Index.remove_files(path)
    before = _foreign_sqlite(path)
    with pytest.raises(ConfigError, match="not a careeros index"):
        Index(data["settings"])
    with pytest.raises(ConfigError, match="not a careeros index"):
        Index.remove_files(path)
    assert path.read_bytes() == before


def test_our_old_schema_is_still_adopted_and_rebuilt(data):
    ix = Index(data["settings"])
    ix.rebuild()
    ix.set_meta("schema_version", "0")
    ix.close()
    ix = Index(data["settings"])
    ix.sync()
    assert ix.get_meta("schema_version") == str(index_mod.SCHEMA_VERSION)
    assert ix.query("SELECT COUNT(*) AS n FROM jobs")[0]["n"] == len(data["jobs"])
    ix.close()


def test_index_path_compared_case_insensitively(data):
    from careeros.config import ConfigError
    from careeros.ui.index import default_path

    s = data["settings"]
    tracker = s.paths["tracker_xlsx"]
    tracker.unlink()                                             # not created yet: still refused
    s.pipeline.setdefault("ui", {})["index_path"] = str(tracker.with_name(tracker.name.upper()))
    with pytest.raises(ConfigError):
        default_path(s)
    assert not tracker.exists()


@pytest.mark.parametrize("probe", ["quick_check", "first_sync"])
def test_damage_inside_a_data_page_is_caught_and_rebuilt(data, monkeypatch, probe):
    ix = Index(data["settings"])
    ix.rebuild()
    root = ix.query("SELECT rootpage FROM sqlite_master WHERE name = 'jobs'")[0]["rootpage"]
    page_size = ix.query("PRAGMA page_size")[0]["page_size"]
    ix.con.execute("PRAGMA wal_checkpoint(TRUNCATE)")
    ix.close()
    path = ix.path
    with path.open("r+b") as f:             # schema and meta pages stay readable; the jobs page does not
        f.seek((root - 1) * page_size)
        f.write(b"\xff" * 100)
    if probe == "first_sync":               # the probe misses it: the first sync must still recover
        monkeypatch.setattr(index_mod, "_classify", lambda p: "ours")
    ix = Index(data["settings"])
    ix.sync()
    assert ix.query("SELECT COUNT(*) AS n FROM jobs")[0]["n"] == len(data["jobs"])
    ix.close()
    assert len(list(path.parent.glob(path.name + ".corrupt-*"))) == 1


def test_a_damaged_index_keeps_its_wal_next_to_the_aside_copy(data):
    path = Index(data["settings"]).path
    Index.remove_files(path)
    path.write_bytes(b"SQLite format 3\x00" + b"damaged page" * 50)
    wal = path.with_name(path.name + "-wal")
    wal.write_bytes(b"wal frames " * 20)
    ix = Index(data["settings"])
    ix.close()
    aside = [p for p in path.parent.glob(path.name + ".corrupt-*") if not p.name.endswith("-wal")]
    assert len(aside) == 1
    assert aside[0].with_name(aside[0].name + "-wal").read_bytes() == b"wal frames " * 20


def test_set_aside_twice_in_the_same_second_gets_a_counter(data, monkeypatch):
    class Frozen(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 9, 24, 15, 0, 0, tzinfo=tz)

    monkeypatch.setattr(index_mod, "datetime", Frozen)
    path = data["settings"].paths["jobs_dir"].parent / "careeros.db"
    path.write_bytes(b"one")
    first = index_mod._set_aside(path)
    path.write_bytes(b"two")
    second = index_mod._set_aside(path)
    assert first.name == "careeros.db.corrupt-20260924-150000"
    assert second.name == "careeros.db.corrupt-20260924-150000-1"
    assert (first.read_bytes(), second.read_bytes()) == (b"one", b"two")


def test_a_pre_detail_tracker_item_is_indexed_with_detail_none(idx, data, monkeypatch):
    legacy = {"ID": "old1", "Type": "other", "What to do": "careeros run: /score-job failed", "Done": "N"}
    monkeypatch.setattr(index_mod, "read_action_items", lambda _p: [legacy])
    monkeypatch.setattr(idx, "get_meta", lambda k: None if k == "tracker_sig" else "x")
    idx.sync()
    assert [(r["id"], r["detail"]) for r in idx.query("SELECT id, detail FROM action_items")] == [("old1", None)]


def test_selected_and_injection_flags_are_indexed(data):
    """REQ-104/109: jobs rows carry `selected` (missing flag = ticked) and the uncleared injection reasons."""
    store = Store(data["settings"])
    jid, other = data["jobs"]["review"], data["jobs"]["found"]
    store.set_selected([jid], False)
    store._write(jid, "flags.json", {**store.load_flags(jid), "injection_suspected": True,
                                     "injection_reasons": ["hidden text", "instruction phrase"]})
    ix = Index(data["settings"])
    ix.rebuild()
    jobs = _jobs(ix)
    assert (jobs[jid]["selected"], jobs[jid]["injection"]) == (0, "hidden text; instruction phrase")
    assert (jobs[other]["selected"], jobs[other]["injection"]) == (1, None)
    store.clear_injection(jid)
    ix.update_jobs([jid])
    assert _jobs(ix)[jid]["injection"] is None
    ix.close()
