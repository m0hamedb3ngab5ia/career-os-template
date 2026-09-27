"""Action Items: optional due date + reason (the Action Items screen's "Add date"), reopen, and the in-place
migration of workbooks made before the Due columns existed."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from careeros.models import ActionItem
from careeros.tracker import Tracker

pytestmark = pytest.mark.unit


def _headers(tr: Tracker) -> list[str]:
    return [c.value for c in load_workbook(tr.path)["Action Items"][1] if c.value]


def test_model_has_optional_due_fields():
    it = ActionItem(what="x")
    assert it.due is None and it.due_reason is None
    it = ActionItem(what="x", due="2026-10-01T18:00:00+00:00", due_reason="posting closes")
    assert it.due.startswith("2026-10-01") and it.due_reason == "posting closes"


def test_model_rejects_a_due_that_is_not_a_date():
    with pytest.raises(ValueError):
        ActionItem(what="x", due="next tuesday")


def test_new_workbook_has_due_columns(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    h = _headers(tr)
    assert "Due" in h and "Due reason" in h


def test_due_round_trips(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    aid = tr.add_action_item("pick a slot", type="review", due="2026-10-01T15:00:00+00:00",
                             due_reason="reply within 48 hours")
    plain = tr.add_action_item("no date")
    items = {i["ID"]: i for i in tr.list_action_items()}
    assert items[aid]["Due"] == "2026-10-01T15:00:00+00:00"
    assert items[aid]["Due reason"] == "reply within 48 hours"
    assert items[plain]["Due"] in (None, "") and items[plain]["Due reason"] in (None, "")


def test_set_action_due_and_clear(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    aid = tr.add_action_item("essay")
    assert tr.set_action_due(aid, "2026-10-03", "application deadline") is True
    it = {i["ID"]: i for i in tr.list_action_items()}[aid]
    assert it["Due"] == "2026-10-03" and it["Due reason"] == "application deadline"
    assert tr.set_action_due(aid, None, None) is True
    it = {i["ID"]: i for i in tr.list_action_items()}[aid]
    assert it["Due"] in (None, "") and it["Due reason"] in (None, "")
    assert tr.set_action_due("nope", "2026-10-03") is False
    with pytest.raises(ValueError):
        tr.set_action_due(aid, "someday")


def test_reopen_action(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    aid = tr.add_action_item("captcha")
    assert tr.mark_action_done(aid) is True
    assert aid not in {i["ID"] for i in tr.list_action_items()}
    assert tr.reopen_action(aid) is True
    it = {i["ID"]: i for i in tr.list_action_items()}[aid]
    assert it["Done"] == "N" and it["DoneDate"] in (None, "")
    assert tr.reopen_action("nope") is False


def test_old_workbook_migrates_due_columns_in_place(tmp_path: Path):
    """A workbook from before Due existed gets Due + Due reason appended; existing cells stay where they were."""
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    wb = load_workbook(tr.path)
    ws = wb["Action Items"]
    hdr = {c.value: c.column for c in ws[1]}
    for name in ("Due reason", "Due"):
        ws.delete_cols([c.column for c in ws[1] if c.value == name][0])
    ws.append(["old1", "2026-01-01 00:00:00", "j1", "Acme", "SWE", "review", "look at it", "", "H", "phone", "N", ""])
    wb.save(tr.path)
    assert "Due" not in _headers(tr)

    aid = tr.add_action_item("new one", due="2026-10-01", due_reason="posting closes")
    h = _headers(tr)
    assert h[-2:] == ["Due", "Due reason"] and h.index("Needs") == hdr["Needs"] - 1
    items = {i["ID"]: i for i in tr.list_action_items()}
    assert items["old1"]["What to do"] == "look at it" and items["old1"]["Needs"] == "phone"
    assert items["old1"]["Due"] in (None, "")
    assert items[aid]["Due"] == "2026-10-01" and items[aid]["Due reason"] == "posting closes"
    # idempotent
    tr.list_action_items()
    assert _headers(tr).count("Due") == 1


def test_set_due_on_old_workbook_migrates_first(tmp_path: Path):
    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    aid = tr.add_action_item("x")
    wb = load_workbook(tr.path)
    ws = wb["Action Items"]
    for name in ("Due reason", "Due"):
        ws.delete_cols([c.column for c in ws[1] if c.value == name][0])
    wb.save(tr.path)
    assert tr.set_action_due(aid, "2026-10-02", "saved form expires") is True
    it = {i["ID"]: i for i in tr.list_action_items()}[aid]
    assert it["Due"] == "2026-10-02" and it["Due reason"] == "saved form expires"


def test_queued_reopen_due_and_add_replay_on_flush(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """Excel holds the workbook: reopen / set due / add-with-due queue to .pending.json and replay in order."""
    import json

    from openpyxl.workbook.workbook import Workbook

    tr = Tracker(path=tmp_path / "t.xlsx")
    tr.init()
    aid = tr.add_action_item("essay")
    assert tr.mark_action_done(aid) is True
    real_save = Workbook.save

    def locked(self, filename):  # noqa: ANN001, ANN202
        raise PermissionError(13, "locked")

    monkeypatch.setattr(Workbook, "save", locked)
    with pytest.warns(UserWarning):
        assert tr.reopen_action(aid) is None
        assert tr.set_action_due(aid, "2026-10-03", "application deadline") is None
        assert tr.add_action_item("captcha", id="a2", due="2026-10-04", due_reason="posting closes") == "a2"
    q = json.loads(tr.pending_path.read_text())
    assert [x["op"] for x in q] == ["reopen_action", "set_action_due", "add_action_item"]
    assert q[2]["payload"]["due"] == "2026-10-04"
    assert aid not in {i["ID"] for i in tr.list_action_items()}  # nothing written yet

    monkeypatch.setattr(Workbook, "save", real_save)
    assert tr.flush_pending() == 3 and tr.pending_count() == 0
    items = {i["ID"]: i for i in tr.list_action_items()}
    assert items[aid]["Done"] == "N" and items[aid]["Due"] == "2026-10-03"
    assert items[aid]["Due reason"] == "application deadline"
    assert items["a2"]["Due"] == "2026-10-04" and items["a2"]["Due reason"] == "posting closes"
