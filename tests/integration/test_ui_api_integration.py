"""The UI server end to end: FastAPI TestClient over a temp repo root with the fictional UI data, the static
frontend fallback, the request guard, and `careeros ui` as a real subprocess (health, then live SSE on a file
change)."""
from __future__ import annotations

import json
import socket
import subprocess
import types
import sys
import time
from datetime import datetime, timezone

import pytest
from conftest import make_temp_root, subprocess_env
from fixtures.ui_data import build_ui_data

pytest.importorskip("fastapi")
httpx = pytest.importorskip("httpx")
from fastapi.testclient import TestClient  # noqa: E402

from careeros.store import Store  # noqa: E402
from careeros.ui.app import create_app  # noqa: E402
from careeros.ui.index import Index  # noqa: E402
from careeros.ui.security import LOOPBACK  # noqa: E402

pytestmark = pytest.mark.integration

NOW = datetime(2026, 9, 24, 15, 0, tzinfo=timezone.utc)


@pytest.fixture
def data(tmp_path):
    return build_ui_data(make_temp_root(tmp_path / "repo"), NOW)


@pytest.fixture
def client(data, tmp_path):
    ix = Index(data["settings"])
    ix.rebuild()
    app = create_app(data["settings"], index=ix, allowed_hosts=LOOPBACK | {"testserver"},
                     static_dir=tmp_path / "no-static", now=lambda: NOW)
    with TestClient(app) as c:
        yield c
    ix.close()


def test_health_and_meta(client):
    h = client.get("/api/health").json()
    assert h["ok"] is True and h["indexed_at"] and h["config_error"] is None
    m = client.get("/api/meta").json()
    assert "needs_review" in m["statuses"] and m["presets"]["recommended"] == "medium"


def test_status(client):
    st = client.get("/api/status").json()
    assert st["tiles"]["needs_you"]["value"] == 3 and st["counts"]["jobs"] == 8
    assert st["tiles"]["applied_week"]["value"] == 1


def test_jobs_list_filters_and_paging(client):
    page = client.get("/api/jobs", params={"limit": 2}).json()
    assert page["total"] == 8 and len(page["items"]) == 2 and page["next_cursor"] == "2"
    applied = client.get("/api/jobs", params=[("status", "applied"), ("status", "interview")]).json()
    assert {j["company"] for j in applied["items"]} == {"Hooli", "Stark Industries"}
    assert client.get("/api/jobs", params={"q": "initech"}).json()["total"] == 1
    remote = client.get("/api/jobs", params={"location": "REMOTE", "sort": "location"}).json()
    assert [j["location"] for j in remote["items"]] == ["Remote"]
    assert client.get("/api/jobs", params={"location": "new york"}).json()["total"] == 7
    tabs = {t["key"]: t["count"] for t in client.get("/api/jobs/tabs", params={"location": "remote"}).json()["tabs"]}
    assert tabs["all"] == 1
    xlsx = client.post("/api/jobs/export", headers={"x-careeros": "1"},
                       json={"tab": "all", "location": "remote", "sort": "-fit", "columns": ["company"]})
    assert xlsx.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook
    rows = list(load_workbook(BytesIO(xlsx.content)).active.iter_rows(values_only=True))
    assert [r[1] for r in rows[1:]] == ["Globex"]
    assert client.get("/api/jobs", params={"sort": "bogus"}).status_code == 400
    assert client.get("/api/jobs", params={"limit": 0}).status_code == 422


def test_jobs_column_filters_facets_tabs_and_export(client):
    r = client.get("/api/jobs", params=[("company", "Hooli"), ("company", "Initech"), ("tab", "all")]).json()
    assert {j["company"] for j in r["items"]} == {"Hooli", "Initech"}
    assert client.get("/api/jobs", params={"fit_min": 85, "fit_max": 90, "tab": "all"}).json()["total"] == 2
    assert client.get("/api/jobs", params={"found_from": "2026-09-19", "found_to": "2026-09-21", "tab": "all"}).json()["total"] == 3
    assert client.get("/api/jobs", params=[("location_in", "Remote"), ("location", "rem")]).json()["total"] == 1
    assert client.get("/api/jobs", params={"fit_min": "high"}).status_code == 422
    assert client.get("/api/jobs", params={"found_from": "yesterday"}).status_code == 422
    assert client.get("/api/jobs", params={"qa_passed": "1", "tab": "all"}).json()["total"] == 4
    assert client.get("/api/jobs", params={"closes_from": "2026-01-01", "closes_to": "2026-12-31",
                                           "tab": "all"}).json()["total"] == 0
    assert client.get("/api/jobs", params={"closes_from": "not-a-date"}).status_code == 422
    f = client.get("/api/jobs/facets", params={"field": "tier", "tab": "all", "tier": "A", "status": "applied"}).json()
    assert f == {"field": "tier", "values": [{"value": "B", "count": 1}]}
    # location facet must drop only the values["location"] filter, not the `location` substring search, so a
    # substring search still narrows the facet counts (regression: field-name kwarg pop used to eat `location`).
    loc_facet = client.get("/api/jobs/facets", params={"field": "location", "tab": "all", "location": "rem"}).json()
    assert loc_facet == {"field": "location", "values": [{"value": "Remote", "count": 1}]}
    assert client.get("/api/jobs/facets", params={"field": "url"}).status_code == 400
    tabs = {t["key"]: t["count"] for t in client.get("/api/jobs/tabs", params={"fit_min": 85}).json()["tabs"]}
    assert tabs["all"] == 3 and tabs["tier_a"] == 2
    xlsx = client.post("/api/jobs/export", headers={"x-careeros": "1"},
                       json={"tab": "all", "company": ["Hooli"], "fit_min": 80, "columns": ["company"]})
    assert xlsx.status_code == 200
    from openpyxl import load_workbook
    import io
    assert load_workbook(io.BytesIO(xlsx.content))["Jobs"].max_row == 2


def test_job_detail(client, data):
    jid = data["jobs"]["review"]
    d = client.get(f"/api/jobs/{jid}").json()
    assert d["job"]["company"] == "Umbrella Labs" and d["safety"]["verdict"] == "review"
    assert client.get("/api/jobs/nope00000000").status_code == 404
    assert client.get("/api/jobs/..%2Fconfig").status_code == 404


def test_guard_refuses_foreign_host_origin_and_headerless_writes(client):
    assert client.get("/api/health", headers={"host": "evil.example"}).status_code == 421
    assert client.get("/api/health", headers={"origin": "https://evil.example"}).status_code == 403
    assert client.post("/api/health").status_code == 403
    assert client.post("/api/health", headers={"x-careeros": "1"}).status_code in (404, 405)  # past the guard


def test_placeholder_page_without_a_built_frontend(client):
    r = client.get("/")
    assert r.status_code == 200 and "careeros ui" in r.text and "/api/health" in r.text
    assert client.get("/api/nope").status_code == 404


def test_static_frontend_and_spa_fallback(data, tmp_path):
    static = tmp_path / "static"
    (static / "assets").mkdir(parents=True)
    (static / "index.html").write_text("<!doctype html><title>app</title>", encoding="utf-8")
    (static / "assets" / "app.js").write_text("console.log(1)", encoding="utf-8")
    (tmp_path / "secret.txt").write_text("no", encoding="utf-8")
    ix = Index(data["settings"])
    ix.rebuild()
    app = create_app(data["settings"], index=ix, allowed_hosts=LOOPBACK | {"testserver"}, static_dir=static)
    with TestClient(app) as c:
        assert c.get("/").text.startswith("<!doctype html>")
        assert c.get("/assets/app.js").text == "console.log(1)"
        assert c.get("/jobs/abc123").text.startswith("<!doctype html>")          # client-side route
        assert c.get("/..%2Fsecret.txt").text != "no"
        assert c.get("/assets/..%2F..%2Fsecret.txt").text != "no"
        assert c.get("/api/unknown").status_code == 404
    ix.close()


def test_config_error_is_reported_not_fatal(data, client):
    (data["settings"].root / "config" / "pipeline.yaml").write_text("ui: {theme: blue}\n", encoding="utf-8")
    client.app.state.ctx.reload_settings()
    h = client.get("/api/health").json()
    assert "ui.theme" in h["config_error"]
    assert client.get("/api/status").status_code == 200        # still serving with the last good settings


# --- real server ------------------------------------------------------------------------------------------------

def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def _wait_health(url: str, proc: subprocess.Popen, timeout: float = 20) -> dict:
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        if proc.poll() is not None:
            raise AssertionError(f"careeros ui exited {proc.returncode}: {proc.stdout.read()}")
        try:
            return httpx.get(url + "/api/health", timeout=1).json()
        except httpx.HTTPError:
            time.sleep(0.2)
    raise AssertionError("careeros ui did not answer /api/health")


@pytest.fixture
def server(data, tmp_path):
    home = tmp_path / "home"
    home.mkdir()
    port = _free_port()
    env = subprocess_env(data["settings"].root, home)
    proc = subprocess.Popen([sys.executable, "-m", "careeros.cli", "ui", "--no-open", "--port", str(port)],
                            env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    url = f"http://127.0.0.1:{port}"
    try:
        yield url, _wait_health(url, proc)
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            proc.kill()


def test_cli_serves_health_and_indexes_on_start(server, data):
    url, health = server
    assert health["ok"] is True
    assert httpx.get(url + "/api/status").json()["counts"]["jobs"] == 8
    assert (data["settings"].paths["jobs_dir"].parent / "careeros.db").exists()


def test_cli_pushes_an_sse_event_when_a_job_file_changes(server, data):
    url, _ = server
    jid = data["jobs"]["queued"]
    events = []
    with httpx.stream("GET", url + "/api/events", timeout=httpx.Timeout(15.0)) as r:
        assert r.headers["content-type"].startswith("text/event-stream")
        lines = r.iter_lines()
        for line in lines:                                   # the hello frame first
            if line.startswith("event: hello"):
                break
        Store(data["settings"]).set_status(jid, "prepared", "docs ready")
        cur = {}
        deadline = time.monotonic() + 15
        for line in lines:
            if line.startswith("event: "):
                cur["event"] = line[7:]
            elif line.startswith("data: "):
                cur["data"] = json.loads(line[6:])
            elif line == "" and cur:
                events.append(cur)
                if cur.get("event") == "changed" and jid in cur["data"]["jobs"]:
                    break
                cur = {}
            if time.monotonic() > deadline:
                break
    assert any(e.get("event") == "changed" and jid in e["data"]["jobs"] for e in events), events
    job = httpx.get(f"{url}/api/jobs/{jid}").json()
    assert job["job"]["status"] == "prepared"


def test_cli_refuses_a_non_loopback_host(data, tmp_path):
    home = tmp_path / "h2"
    home.mkdir()
    r = subprocess.run([sys.executable, "-m", "careeros.cli", "ui", "--no-open", "--host", "0.0.0.0"],
                       env=subprocess_env(data["settings"].root, home), capture_output=True, text=True, timeout=30)
    assert r.returncode == 2 and "127.0.0.1" in r.stderr


def test_cli_reindex_flag_rebuilds(data, tmp_path):
    ix = Index(data["settings"])
    ix.rebuild()
    ix.query("UPDATE jobs SET company = 'Stale Co'")      # drifted rows a plain sync would skip (same files)
    ix.close()
    home = tmp_path / "h3"
    home.mkdir()
    port = _free_port()
    proc = subprocess.Popen([sys.executable, "-m", "careeros.cli", "ui", "--no-open", "--reindex", "--port",
                             str(port)], env=subprocess_env(data["settings"].root, home), stdout=subprocess.PIPE,
                            stderr=subprocess.STDOUT, text=True)
    try:
        _wait_health(f"http://127.0.0.1:{port}", proc)
        assert httpx.get(f"http://127.0.0.1:{port}/api/status").json()["counts"]["jobs"] == 8
        assert httpx.get(f"http://127.0.0.1:{port}/api/jobs", params={"q": "stale"}).json()["total"] == 0
    finally:
        proc.terminate()
        proc.wait(timeout=10)


# --- watcher threads + settings reload -------------------------------------------------------------------------

class _Broker:
    def __init__(self):
        import threading

        self.sent, self.cond = [], threading.Condition()

    def publish(self, event, payload):
        with self.cond:
            self.sent.append((event, payload))
            self.cond.notify_all()
        return len(self.sent)

    def wait_for(self, pred, timeout=15):
        with self.cond:
            return self.cond.wait_for(lambda: any(pred(e, p) for e, p in self.sent), timeout)


def test_watcher_watches_a_tracker_outside_the_data_dirs(tmp_path):
    import yaml

    from careeros.tracker import Tracker
    from careeros.ui.watch import Watcher

    root = make_temp_root(tmp_path / "repo")
    outside = tmp_path / "Desktop"
    outside.mkdir()
    pl = yaml.safe_load((root / "config" / "pipeline.yaml").read_text())
    pl["paths"]["tracker_xlsx"] = str(outside / "JobTracker.xlsx")
    (root / "config" / "pipeline.yaml").write_text(yaml.safe_dump(pl, sort_keys=False))
    data = build_ui_data(root, NOW)
    s = data["settings"]
    assert s.paths["tracker_xlsx"].parent == outside
    ix = Index(s)
    ix.rebuild()
    b = _Broker()
    w = Watcher(s, ix, b, debounce_ms=100)
    w.start()
    try:
        assert len(w._threads) == 2
        time.sleep(0.5)                                     # let both watchers arm
        Tracker(settings=s).mark_action_done(data["actions"]["high"])
        assert b.wait_for(lambda e, p: e == "changed" and p["actions"] is True), b.sent
    finally:
        w.stop()
        ix.close()
    assert ix.path.exists()


def test_reload_keeps_settings_when_paths_change(data, client):
    import yaml

    ctx = client.app.state.ctx
    old = ctx.settings
    cfg = data["settings"].root / "config" / "pipeline.yaml"
    pl = yaml.safe_load(cfg.read_text())
    pl["paths"]["jobs_dir"] = "elsewhere/jobs"
    cfg.write_text(yaml.safe_dump(pl, sort_keys=False))
    ctx.reload_settings()
    assert ctx.settings is old and "restart careeros ui" in ctx.config_error
    pl["paths"]["jobs_dir"] = "data/jobs"
    pl["ui"] = {"index_path": "other.db"}
    cfg.write_text(yaml.safe_dump(pl, sort_keys=False))
    ctx.reload_settings()
    assert ctx.settings is old and "restart careeros ui" in ctx.config_error


def test_reload_accepts_a_harmless_change(data, client):
    import yaml

    ctx = client.app.state.ctx
    cfg = data["settings"].root / "config" / "pipeline.yaml"
    pl = yaml.safe_load(cfg.read_text())
    pl["ui"] = {"theme": "dark"}
    cfg.write_text(yaml.safe_dump(pl, sort_keys=False))
    ctx.reload_settings()
    assert ctx.config_error is None and client.get("/api/meta").json()["ui"]["theme"] == "dark"


@pytest.mark.parametrize("block", [
    {"schedule": {"scout": "0 7 * * *"}},
    {"advisor": {"advise_after_days": -1}},
])
def test_reload_rejects_a_broken_schedule_or_advisor(data, client, block):
    import yaml

    ctx = client.app.state.ctx
    old = ctx.settings
    cfg = data["settings"].root / "config" / "pipeline.yaml"
    pl = yaml.safe_load(cfg.read_text())
    pl.update(block)
    cfg.write_text(yaml.safe_dump(pl, sort_keys=False))
    ctx.reload_settings()
    assert ctx.settings is old and ctx.config_error and ("schedule" in ctx.config_error or
                                                          "advis" in ctx.config_error)


@pytest.mark.parametrize("change", [{"watch_debounce_ms": 900}, {"port": 9100}, {"host": "localhost"}])
def test_reload_keeps_settings_when_server_keys_change(data, client, change):
    import yaml

    ctx = client.app.state.ctx
    old = ctx.settings
    cfg = data["settings"].root / "config" / "pipeline.yaml"
    pl = yaml.safe_load(cfg.read_text())
    pl["ui"] = change
    cfg.write_text(yaml.safe_dump(pl, sort_keys=False))
    ctx.reload_settings()
    assert ctx.settings is old and "restart careeros ui" in ctx.config_error


def test_reload_refuses_an_index_path_on_the_tracker(data, client):
    import yaml

    ctx = client.app.state.ctx
    old = ctx.settings
    cfg = data["settings"].root / "config" / "pipeline.yaml"
    pl = yaml.safe_load(cfg.read_text())
    pl["ui"] = {"index_path": str(data["settings"].paths["tracker_xlsx"])}
    cfg.write_text(yaml.safe_dump(pl, sort_keys=False))
    ctx.reload_settings()
    assert ctx.settings is old and "index_path" in ctx.config_error


def test_reload_refuses_a_broken_volume_block(data, client):
    import yaml

    ctx = client.app.state.ctx
    old = ctx.settings
    cfg = data["settings"].root / "config" / "targets.yaml"
    tg = yaml.safe_load(cfg.read_text()) or {}
    tg["volume"] = {"max_applications_per_day": 0}
    cfg.write_text(yaml.safe_dump(tg, sort_keys=False))
    ctx.reload_settings()
    assert ctx.settings is old and "volume" in ctx.config_error
    assert client.get("/api/status").status_code == 200


def test_cli_refuses_to_replace_a_foreign_file_at_the_index_path(data, tmp_path):
    db = data["settings"].paths["jobs_dir"].parent / "careeros.db"
    db.write_bytes(b"not a database")
    home = tmp_path / "h4"
    home.mkdir()
    r = subprocess.run([sys.executable, "-m", "careeros.cli", "ui", "--no-open", "--reindex"],
                       env=subprocess_env(data["settings"].root, home), capture_output=True, text=True, timeout=30)
    assert r.returncode == 1 and "not a careeros index" in r.stderr
    assert db.read_bytes() == b"not a database"


def test_cli_reindex_refuses_another_apps_sqlite_file(data, tmp_path):
    import sqlite3

    db = data["settings"].paths["jobs_dir"].parent / "careeros.db"
    con = sqlite3.connect(db)
    con.execute("CREATE TABLE jobs (id INTEGER)")
    con.commit()
    con.close()
    before = db.read_bytes()
    home = tmp_path / "h5"
    home.mkdir()
    r = subprocess.run([sys.executable, "-m", "careeros.cli", "ui", "--no-open", "--reindex"],
                       env=subprocess_env(data["settings"].root, home), capture_output=True, text=True, timeout=30)
    assert r.returncode == 1 and "not a careeros index" in r.stderr
    assert db.read_bytes() == before


def test_application_status_open_and_confirmation(client, data, monkeypatch):
    """Staged form reachability: dead tab -> needs refill + a detached fill; live tab -> focus; confirmation -> applied."""
    from careeros.apply import browser
    from careeros.ui.services import job_actions

    jid = data["jobs"]["review"]
    jdir = Store(data["settings"]).job_dir(jid)
    assert client.get(f"/api/jobs/{jid}/application").json()["tab"] == "none"
    browser.save_record(jdir, tab_id="T1", url="https://job-boards.example/embed", cdp=browser.DEFAULT_CDP)
    live: list[dict] = []

    def get(url):
        if "/json/activate/" in url:
            raise ValueError("Target activated")
        if not live:
            raise OSError("refused")
        return live

    monkeypatch.setattr(browser, "_get", get)
    spawned: list[list[str]] = []
    monkeypatch.setattr(job_actions.subprocess, "Popen", lambda argv, **k: spawned.append(argv) or types.SimpleNamespace(pid=999_999))
    from careeros.apply import gh_fill

    def missing():
        raise gh_fill.MissingPlaywright("playwright is not installed")

    monkeypatch.setattr(gh_fill, "preflight", missing)
    h = {"x-careeros": "1"}
    assert client.get(f"/api/jobs/{jid}/application").json()["tab"] == "needs_refill"
    r = client.post(f"/api/jobs/{jid}/application/open", headers=h)
    assert r.status_code == 409 and "playwright is not installed" in r.json()["detail"] and not spawned
    monkeypatch.setattr(gh_fill, "preflight", lambda: None)
    r = client.post(f"/api/jobs/{jid}/application/open", headers=h).json()
    assert r["action"] == "filling" and "apply plan" in spawned[0][-1] and "apply fill" in spawned[0][-1]
    assert browser.FILL_EXIT in spawned[0][-1]
    (jdir / "application.log").write_text(f"boom: chromium missing\n{browser.FILL_EXIT}1\n")
    assert client.get(f"/api/jobs/{jid}/application").json()["fill_error"] == "boom: chromium missing"
    live.append({"id": "T1", "type": "page", "url": "https://job-boards.example/embed"})
    assert client.get(f"/api/jobs/{jid}/application").json()["tab"] == "open"
    assert client.post(f"/api/jobs/{jid}/application/open", headers=h).json() == {"action": "focused"}
    assert len(spawned) == 1
    # Refill skips the live tab and fills again; a staged fill's leftover fields show as fields_left, not an error
    assert client.post(f"/api/jobs/{jid}/application/open", headers=h, json={"refill": True}).json()["action"] == "filling"
    assert len(spawned) == 2
    (jdir / "fill_summary.json").write_text('{"filled": 3, "failed": [{"field_id": "q1", "label": "Q1", "error": "x"}], "skipped": ["Q2"]}')
    assert client.get(f"/api/jobs/{jid}/application").json()["fields_left"] == ["Q1", "Q2"]
    live[0]["url"] = "https://job-boards.example/acme/jobs/1/confirmation"
    assert client.get(f"/api/jobs/{jid}/application").json()["marked_applied"] is True
    assert Store(data["settings"]).get_status(jid) == "applied"


def test_application_open_refuses_when_own_chrome_not_connected(client, data, monkeypatch):
    """paths.apply_cdp (your own Chrome) not answering: a plain 409 to connect and retry, no fill spawned, and the
    recorded tab is kept so the retry focuses it once Chrome is back."""
    from careeros.apply import browser, gh_fill
    from careeros.ui.services import job_actions

    jid = data["jobs"]["review"]
    jdir = Store(data["settings"]).job_dir(jid)
    cdp = "http://127.0.0.1:9222"
    monkeypatch.setitem(data["settings"].paths, "apply_cdp", cdp)
    browser.save_record(jdir, tab_id="T1", url="https://job-boards.example/embed", cdp=cdp)
    live: list[dict] = []

    def get(url):
        if "/json/activate/" in url:
            raise ValueError("Target activated")
        if not live:
            raise OSError("refused")
        return live

    monkeypatch.setattr(browser, "_get", get)
    monkeypatch.setattr(gh_fill, "preflight", lambda: None)
    spawned: list = []
    monkeypatch.setattr(job_actions.subprocess, "Popen", lambda argv, **k: spawned.append(argv) or types.SimpleNamespace(pid=999_999))
    h = {"x-careeros": "1"}
    r = client.post(f"/api/jobs/{jid}/application/open", headers=h)
    assert r.status_code == 409 and r.json()["detail"].startswith("Chrome not connected") and not spawned
    assert browser.record(jdir)["tab_id"] == "T1"
    live.append({"id": "T1", "type": "page", "url": "https://job-boards.example/embed"})
    assert client.post(f"/api/jobs/{jid}/application/open", headers=h).json() == {"action": "focused"}


def test_jobs_select_endpoint_sets_flag(client, data):
    """REQ-104: POST /api/jobs/select ticks/unticks jobs (flags.json `selected`); unknown ids -> 404, nothing set."""
    from careeros.store import Store

    store = Store(data["settings"])
    ids = sorted(store.iter_job_ids())[:2]
    r = client.post("/api/jobs/select", headers={"x-careeros": "1"}, json={"ids": ids, "selected": False})
    assert r.status_code == 200, r.text
    assert r.json() == {"ids": ids, "selected": False}
    assert not any(store.is_selected(j) for j in ids)
    r = client.post("/api/jobs/select", headers={"x-careeros": "1"}, json={"ids": [ids[0], "nope"], "selected": True})
    assert r.status_code == 404 and "nope" in r.text
    assert not store.is_selected(ids[0])
    assert client.post("/api/jobs/select", headers={"x-careeros": "1"}, json={"ids": ids, "selected": True}).status_code == 200
    assert all(store.is_selected(j) for j in ids)


def _write_plan(data, jid):
    from careeros.store import Store

    plan = {"job_id": jid, "ats": "greenhouse", "board": "acme", "ats_job_id": "1", "files": {}, "fields": [
        {"field_id": "email", "label": "Email", "type": "text", "value": "a@example.com", "source": "profile",
         "needs_review": False, "required": True},
        {"field_id": "q1", "label": "What is your notice period?", "type": "text", "value": None,
         "source": "unanswered", "needs_review": True, "required": True},
        {"field_id": "q2", "label": "Favourite editor", "type": "text", "value": None, "source": "unanswered",
         "needs_review": True, "required": False},
        {"field_id": "q3", "label": "Will you need sponsorship?", "type": "select", "value": None,
         "source": "pause:legal", "needs_review": True, "required": True, "options": ["Yes", "No"]}]}
    Store(data["settings"])._write(jid, "fill_plan.json", plan)
    return plan


def test_fill_plan_preview_edit_and_save_to_profile(client, data):
    """REQ-105/106, E2E-008-01/02: the plan table, edits land in fill_plan.json, save-to-profile learns the answer
    and the next plan auto-fills it; required unanswered blocks the fill, optional ones may be skipped."""
    import yaml

    from careeros.apply.gh_schema import build_plan

    jid = sorted(Store(data["settings"]).iter_job_ids())[0]
    h = {"x-careeros": "1"}
    assert client.get(f"/api/jobs/{jid}/fill-plan").json() == {"plan": None, "problems": []}
    _write_plan(data, jid)
    r = client.get(f"/api/jobs/{jid}/fill-plan").json()
    assert [f["field_id"] for f in r["plan"]["fields"]] == ["email", "q1", "q2", "q3"]
    assert r["problems"] == ["unanswered (pause:legal): Will you need sponsorship?",
                             "needs input (required): What is your notice period?"]
    url = f"/api/jobs/{jid}/fill-plan/fields"
    assert client.post(f"{url}/q1", headers=h, json={"skip": True}).status_code == 400  # required: no skip
    assert client.post(f"{url}/q3", headers=h, json={"value": "Maybe", "save": False}).status_code == 400  # not an option
    assert client.post(f"{url}/nope", headers=h, json={"value": "x"}).status_code == 404
    r = client.post(f"{url}/q2", headers=h, json={"skip": True}).json()
    assert r["field"]["skipped"] is True and r["saved"] is False
    sa = data["settings"].paths["standard_answers"]
    n = len(yaml.safe_load(open(sa))["answers"])
    r = client.post(f"{url}/q3", headers=h, json={"value": "No", "save": False}).json()
    assert r["field"]["value"] == "No" and r["field"]["source"] == "user" and r["saved"] is False
    r = client.post(f"{url}/q1", headers=h, json={"value": "4 weeks", "save": True}).json()
    assert r["saved"] is True and r["problems"] == []
    plan = json.loads((Store(data["settings"]).job_dir(jid) / "fill_plan.json").read_text())
    assert {f["field_id"]: f["value"] for f in plan["fields"]}["q1"] == "4 weeks"
    answers = yaml.safe_load(open(sa))["answers"]
    assert len(answers) == n + 1 and answers[-1]["answer"] == "4 weeks"  # save off (q3): job only
    nxt = build_plan([{"field_id": "q9", "label": "What is your notice period?", "type": "text", "options": [],
                       "required": True}], profile={}, answers_path=sa, files={})
    assert nxt["fields"][0]["value"] == "4 weeks"


def test_application_open_refuses_required_needs_input(client, data, monkeypatch):
    """REQ-106: a plan with a required unanswered field is not filled; no fill is spawned."""
    from careeros.apply import gh_fill
    from careeros.ui.services import job_actions

    jid = sorted(Store(data["settings"]).iter_job_ids())[0]
    _write_plan(data, jid)
    monkeypatch.setattr(gh_fill, "preflight", lambda: None)
    spawned: list = []
    monkeypatch.setattr(job_actions.subprocess, "Popen", lambda argv, **k: spawned.append(argv) or types.SimpleNamespace(pid=1))
    r = client.post(f"/api/jobs/{jid}/application/open", headers={"x-careeros": "1"})
    assert r.status_code == 400 and "needs input (required)" in r.json()["detail"] and not spawned
