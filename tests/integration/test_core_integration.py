"""Core boundaries: recorded HTTP -> scout -> data/jobs/<id>/ -> CLI subprocess -> JobTracker.xlsx."""
from __future__ import annotations

import json
import subprocess
from pathlib import Path

import pytest
import yaml
from conftest import FIXTURES, PY, load_fixture, subprocess_env
from openpyxl import load_workbook

import careeros.scout.base as base
from careeros.config import Settings
from careeros.models import Posting
from careeros.scout import run_scout
from careeros.store import Store

pytestmark = pytest.mark.integration

# Recorded board responses keyed by the exact API URL each adapter requests.
RECORDED = {
    "https://boards-api.greenhouse.io/v1/boards/acme/jobs": (200, "greenhouse.json"),
    "https://api.lever.co/v0/postings/leverco": (200, "lever.json"),
    "https://api.ashbyhq.com/posting-api/job-board/ashbyco": (200, "ashby.json"),
    "https://boards-api.greenhouse.io/v1/boards/nosuchslug/jobs": (404, None),
}


class _Resp:
    def __init__(self, status: int, payload):
        self.status_code, self._payload = status, payload

    def raise_for_status(self):
        if self.status_code >= 400:
            import requests

            raise requests.HTTPError(str(self.status_code))

    def json(self):
        return self._payload


@pytest.fixture
def recorded_http(monkeypatch):
    requested: list[str] = []

    def fake_get(url, params=None, headers=None, timeout=None):
        requested.append(url)
        if url not in RECORDED:
            raise AssertionError(f"unexpected network call: {url}")
        status, name = RECORDED[url]
        return _Resp(status, load_fixture(name) if name else None)

    monkeypatch.setattr(base.requests, "get", fake_get)
    monkeypatch.setattr(base.time, "sleep", lambda s: None)
    return requested


def _cli(root: Path, home: Path, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run([PY, "-m", "careeros.cli", "--root", str(root), *args], capture_output=True, text=True,
                          env=subprocess_env(root, home), timeout=120)


@pytest.fixture
def home(tmp_path: Path) -> Path:
    h = tmp_path / "home"
    h.mkdir()
    return h


def test_scout_to_disk_to_tracker_sync(temp_root: Path, home: Path, recorded_http):
    companies = yaml.safe_load((temp_root / "config" / "companies.yaml").read_text())
    companies["boards"].append({"company": "Ghost Co", "ats": "greenhouse", "slug": "nosuchslug"})
    (temp_root / "config" / "companies.yaml").write_text(yaml.safe_dump(companies))
    targets = yaml.safe_load((temp_root / "config" / "targets.yaml").read_text())
    targets["location"]["blocked_countries"] = ["DE"]  # the example blocks nothing; exercise the filter
    # recorded payloads carry fixed dates; keep the stale check out of this test (test_ghost_integration covers it)
    targets.setdefault("safety", {})["ghost"] = {"stale_flag_days": 100000, "stale_skip_days": 100000}
    (temp_root / "config" / "targets.yaml").write_text(yaml.safe_dump(targets))

    s = Settings.load(temp_root)
    store = Store(s)
    logs: list[str] = []
    summary = run_scout(s, store, log=logs.append)

    # bad slug is reported, not raised; custom board skipped
    assert [b.company for b in summary.bad_slugs] == ["Ghost Co"]
    assert {b.company: b.status for b in summary.boards}["Custom Co"] == "skipped"
    assert any("404 bad slug" in line for line in logs)
    assert set(recorded_http) == set(RECORDED)

    # recorded payloads: GH backend SWE (kept) + PM (title), Lever DE (kept) + Berlin (location), Ashby FS (kept)
    t = summary.totals
    assert (t["fetched"], t["stored"], t["filtered_title"], t["filtered_location"]) == (5, 3, 1, 1)
    jobs_dir = temp_root / "data" / "jobs"
    posting_files = sorted(jobs_dir.glob("*/posting.json"))
    assert len(posting_files) == 3
    titles = {json.loads(p.read_text())["title"] for p in posting_files}
    assert titles == {"Software Engineer, Backend", "Data Engineer", "Full Stack Engineer"}
    for p in posting_files:
        assert json.loads((p.parent / "status.json").read_text())["status"] == "found"
        assert "found via" in (p.parent / "log.md").read_text()
    assert len(json.loads((temp_root / "data" / "seen.json").read_text())) == 5

    # second run: nothing new, nothing re-stored
    again = run_scout(Settings.load(temp_root), Store(Settings.load(temp_root)), log=lambda m: None)
    assert again.totals["new"] == 0 and again.totals["stored"] == 0

    r = _cli(temp_root, home, "tracker", "sync")
    assert r.returncode == 0, r.stderr
    assert "synced 3 jobs" in r.stdout
    xlsx = temp_root / "JobTracker.xlsx"
    ws = load_workbook(xlsx)["Jobs"]
    hdr = {c.value: c.column - 1 for c in ws[1]}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr["JobID"]]]
    assert {r[hdr["Role"]] for r in rows} == titles
    assert {r[hdr["Status"]] for r in rows} == {"found"}
    assert {r[hdr["Company"]] for r in rows} == {"Acme", "Lever Co", "Ashby Co"}
    by_role = {r[hdr["Role"]]: r for r in rows}
    assert by_role["Data Engineer"][hdr["Salary"]] == "140,000-180,000 USD"
    assert not list(home.rglob("*.xlsx")), "nothing may be written under HOME"


def _seed_job(root: Path, jid: str = "job000000001") -> str:
    Store(Settings.load(root)).save_posting(Posting(job_id=jid, company="Initech", title="Backend Engineer",
                                                    location="Boston, MA", ats="greenhouse", url="https://x/1"))
    return jid


def test_cli_subprocess_roundtrip(temp_root: Path, home: Path):
    jid = _seed_job(temp_root)

    r = _cli(temp_root, home, "tracker", "init")
    assert r.returncode == 0 and "created" in r.stdout
    assert (temp_root / "JobTracker.xlsx").exists()
    assert _cli(temp_root, home, "tracker", "init").stdout.startswith("exists")

    r1 = _cli(temp_root, home, "action", "add", "review tier A", "--type", "review", "--job", jid,
              "--needs", "laptop", "--dedupe")
    assert r1.returncode == 0 and "added (review/M/laptop)" in r1.stdout
    r2 = _cli(temp_root, home, "action", "add", "review again", "--type", "review", "--job", jid, "--dedupe")
    assert r2.returncode == 0 and "already open" in r2.stdout
    listed = _cli(temp_root, home, "action", "list").stdout.strip().splitlines()
    assert len(listed) == 1 and "Initech" in listed[0] and "laptop" in listed[0]

    r = _cli(temp_root, home, "job", "status", jid, "queued", "--note", "qa pass")
    assert r.returncode == 0, r.stderr
    assert json.loads((temp_root / "data" / "jobs" / jid / "status.json").read_text())["status"] == "queued"
    ws = load_workbook(temp_root / "JobTracker.xlsx")["Jobs"]
    hdr = {c.value: c.column - 1 for c in ws[1]}
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr["JobID"]] == jid)
    assert row[hdr["Status"]] == "queued" and "qa pass" in row[hdr["Notes"]]
    bad = _cli(temp_root, home, "job", "status", "nope", "queued")
    assert bad.returncode == 1 and "not found" in bad.stderr
    assert _cli(temp_root, home, "job", "status", jid, "bogus").returncode == 2  # argparse choices

    r = _cli(temp_root, home, "jobs", "list", "--json")
    jobs = json.loads(r.stdout)
    assert [(j["job_id"], j["status"], j["company"]) for j in jobs] == [(jid, "queued", "Initech")]
    assert json.loads(_cli(temp_root, home, "jobs", "list", "--json", "--status", "applied").stdout) == []

    r = _cli(temp_root, home, "stats")
    assert r.returncode == 0
    assert "jobs in data: 1" in r.stdout and "queued=1" in r.stdout and "open actions=1" in r.stdout

    show = _cli(temp_root, home, "job", "show", jid)
    assert "Initech — Backend Engineer" in show.stdout and "status -> queued: qa pass" in show.stdout
    assert not any(home.iterdir()), "CLI must not write under HOME"


def test_cli_tracker_upsert_and_applied_count(temp_root: Path, home: Path):
    """The exact calls the apply-job skill makes, as subprocesses on a temp root."""
    assert _cli(temp_root, home, "tracker", "applied-count", "Initech", "--days", "90").stdout.strip() == "0"
    assert not (temp_root / "JobTracker.xlsx").exists()
    jid = _seed_job(temp_root)
    r = _cli(temp_root, home, "tracker", "upsert", jid, "--field", "DateApplied=today", "--field", "ATS=greenhouse",
             "--field", "ResumeVersion=swe_backend-v1", "--field", "Status=applied", "--field", "Company=Initech")
    assert r.returncode == 0, r.stderr
    assert f"{jid}: created" in r.stdout
    ws = load_workbook(temp_root / "JobTracker.xlsx")["Jobs"]
    hdr = {c.value: c.column - 1 for c in ws[1]}
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr["JobID"]] == jid)
    assert (row[hdr["Status"]], row[hdr["ATS"]], row[hdr["ResumeVersion"]]) == ("applied", "greenhouse", "swe_backend-v1")
    assert row[hdr["DateApplied"]] and row[hdr["Folder"]] == "open"
    r = _cli(temp_root, home, "tracker", "applied-count", "initech inc.", "--days", "90")
    assert r.returncode == 0 and r.stdout.strip() == "1"
    bad = _cli(temp_root, home, "tracker", "upsert", jid, "--field", "Status=hired")
    assert bad.returncode == 2 and "invalid status" in bad.stderr
    assert not any(home.iterdir()), "CLI must not write under HOME"


def test_cli_stats_without_tracker(temp_root: Path, home: Path):
    r = _cli(temp_root, home, "stats")
    assert r.returncode == 0 and "tracker: not created" in r.stdout


# --- YAML schema smoke ---------------------------------------------------------

REPO = FIXTURES.parents[1]


# Only the shipped examples: tests never read the personal config/ + profile/ (gitignored, often
# symlinks into a private repo outside this checkout, absent in CI).
def _roots():
    yield pytest.param(REPO / "examples", id="examples")


def _strict(path: Path) -> dict:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))  # raises on invalid YAML (no warning swallowing)
    assert isinstance(data, dict), f"{path} must be a mapping"
    return data


@pytest.mark.parametrize("root", list(_roots()))
def test_yaml_schema_smoke(root: Path):
    """The same required-key rules `careeros doctor` applies to a candidate's own files (doctor.schema_problems),
    plus what only the shipped example must satisfy."""
    from careeros.doctor import STANDARD_KEYS, schema_problems

    cfg = {n: _strict(root / "config" / f"{n}.yaml") for n in ("targets", "categories", "companies", "qa", "pipeline")}
    prof = {n: _strict(root / "profile" / f"{n}.yaml") for n in ("master", "standard_answers", "confidential_terms")}
    assert schema_problems(cfg, prof) == []

    # example-only expectations
    t, co = cfg["targets"], cfg["companies"]
    assert t["location"].get("blocked_countries") == [] and co.get("already_applied") == []  # neutral example
    keys = [a["key"] for a in prof["standard_answers"]["answers"]]
    assert set(keys) == set(STANDARD_KEYS), set(keys) ^ set(STANDARD_KEYS)  # exactly the full set, nothing extra

    s = Settings.load(root)  # and the loader agrees
    assert s.boards and s.title_keywords()


def test_cli_upsert_status_survives_tracker_sync(temp_root: Path, home: Path):
    """apply-job's `tracker upsert --field Status=applied` must not be reverted by the next `tracker sync`."""
    jid = _seed_job(temp_root)
    assert _cli(temp_root, home, "tracker", "upsert", jid, "--field", "Status=applied",
                "--field", "DateApplied=today").returncode == 0
    assert json.loads((temp_root / "data" / "jobs" / jid / "status.json").read_text())["status"] == "applied"
    r = _cli(temp_root, home, "tracker", "sync")
    assert r.returncode == 0, r.stderr
    ws = load_workbook(temp_root / "JobTracker.xlsx")["Jobs"]
    hdr = {c.value: c.column - 1 for c in ws[1]}
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr["JobID"]] == jid)
    assert row[hdr["Status"]] == "applied"
