"""The Today, Jobs and Job detail API end to end over a temp repo root with the fictional UI data: FastAPI
TestClient drives each endpoint and the test checks the files the domain code wrote (status.json, the tracker,
the safety registry) and the refusals (no header, traversal, unknown ids, busy / paused runs)."""
from __future__ import annotations

import io
import os
from datetime import datetime, timezone
from typing import Any

import pytest
from conftest import make_temp_root
from fixtures.ui_data import build_ui_data

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from careeros.runs.store import RunStore  # noqa: E402
from careeros.safety import registry  # noqa: E402
from careeros.store import Store  # noqa: E402
from careeros.tracker import Tracker  # noqa: E402
from careeros.ui.app import create_app  # noqa: E402
from careeros.ui.index import Index  # noqa: E402
from careeros.ui.security import LOOPBACK  # noqa: E402
from careeros.ui.services import desktop  # noqa: E402
from careeros.ui.services.runs import RunControl  # noqa: E402

pytestmark = pytest.mark.integration

NOW = datetime(2026, 9, 24, 15, 0, tzinfo=timezone.utc)
H = {"X-CareerOS": "1"}


class FakeProc:
    pid = 4242


class FakePopen:
    def __init__(self) -> None:
        self.calls: list[list[str]] = []

    def __call__(self, argv: list[str], **kw: Any) -> FakeProc:
        self.calls.append(list(argv))
        return FakeProc()


@pytest.fixture
def data(tmp_path):
    return build_ui_data(make_temp_root(tmp_path / "repo"), NOW)


@pytest.fixture
def popen():
    return FakePopen()


@pytest.fixture
def opened(monkeypatch):
    calls: list[list[str]] = []
    monkeypatch.setattr(desktop, "_run", lambda argv: calls.append(list(argv)))
    monkeypatch.setattr(desktop, "_platform", lambda: "darwin")
    return calls


@pytest.fixture
def client(data, tmp_path, popen):
    ix = Index(data["settings"])
    ix.rebuild()
    app = create_app(data["settings"], index=ix, allowed_hosts=LOOPBACK | {"testserver"},
                     static_dir=tmp_path / "no-static", now=lambda: NOW)
    app.state.ctx.run_control = RunControl(data["settings"], popen=popen, python="python", now=lambda: NOW,
                                           pid_alive=lambda pid: False)
    with TestClient(app) as c:
        c.ix = ix  # type: ignore[attr-defined]
        yield c
    ix.close()


# --- Today -------------------------------------------------------------------------------------------------------

def test_today_lists_open_items_and_the_prepare_queue(client):
    t = client.get("/api/today").json()
    assert [a["company"] for a in t["actions"]] == ["Umbrella Labs", "Initech", "Stark Industries"]
    assert t["actions"][0]["due"] is None and t["prepare_queue"]["error"] is None


def test_mark_done_and_undo(client, data):
    aid = data["actions"]["medium"]
    assert client.post(f"/api/today/actions/{aid}/done").status_code == 403       # no X-CareerOS header
    assert client.post(f"/api/today/actions/{aid}/done", headers=H).json() == {"ok": True, "queued": False}
    open_ids = {i["ID"] for i in Tracker(settings=data["settings"]).list_action_items()}
    assert aid not in open_ids
    client.ix.update_tracker()
    assert aid not in {a["id"] for a in client.get("/api/today").json()["actions"]}
    assert client.post(f"/api/today/actions/{aid}/reopen", headers=H).json()["ok"] is True
    assert aid in {i["ID"] for i in Tracker(settings=data["settings"]).list_action_items()}
    assert client.post("/api/today/actions/nope/done", headers=H).status_code == 404


def test_status_has_the_response_breakdown(client):
    rr = client.get("/api/status").json()["tiles"]["response_rate"]
    assert {b["status"] for b in rr["breakdown"]} >= {"interview", "rejected", "no_reply"}


def test_run_controls(client, data, popen):
    r = client.post("/api/runs/steps/scout", headers=H)
    assert r.status_code == 200 and r.json()["kind"] == "scout" and "scout" in popen.calls[-1]
    r = client.post("/api/runs/batches/prepare", json={"preset": "medium"}, headers=H)
    assert r.status_code == 200 and popen.calls[-1][-5:] == ["run", "prepare", "--preset", "medium", "--json"]
    assert client.post("/api/runs/batches/prepare", json={"preset": "huge"}, headers=H).status_code == 400
    rs = RunStore(data["settings"])
    rs.set_pause(None, "testing", NOW)
    r = client.post("/api/runs/batches/prepare", json={}, headers=H)
    assert r.status_code == 409 and "paused" in r.json()["detail"]
    assert client.post("/api/runs/resume", headers=H).json() == {"resumed": True}
    assert rs.pause_state(NOW) is None


def test_catch_up_run_and_dismiss(client, data, popen):
    rs = RunStore(data["settings"])
    assert client.post("/api/runs/catch-up", json={}, headers=H).json() == {"started": False, "pending": False}
    (rs.dir / "catch_up.json").write_text('{"kinds": {"score": {"slots": 2, "first_missed": null}}}')
    r = client.post("/api/runs/catch-up", json={"dismiss": False}, headers=H).json()
    assert r["pending"] is True and r["kinds"] == ["score"] and "catch-up" in popen.calls[-1]
    r = client.post("/api/runs/catch-up", json={"dismiss": True}, headers=H).json()
    assert r["dismissed"] is True and not (rs.dir / "catch_up.json").exists()


# --- Jobs --------------------------------------------------------------------------------------------------------

def test_tabs_and_tab_filter(client):
    tabs = {t["key"]: t["count"] for t in client.get("/api/jobs/tabs").json()["tabs"]}
    assert tabs == {"active": 6, "review": 1, "applied": 2, "tier_a": 2, "all": 8}
    page = client.get("/api/jobs", params={"tab": "review"}).json()
    assert [j["company"] for j in page["items"]] == ["Umbrella Labs"]
    assert page["items"][0]["next_action"] == "Review and submit"
    assert client.get("/api/jobs", params={"tab": "nope"}).status_code == 400


def test_export_streams_an_xlsx_without_touching_the_tracker(client, data):
    tr = Tracker(settings=data["settings"]).path
    before = tr.stat().st_mtime_ns
    r = client.post("/api/jobs/export", json={"job_ids": [data["jobs"]["applied"]], "columns": ["company", "status"]},
                    headers=H)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows == [("JobID", "Company", "Status"), (data["jobs"]["applied"], "Hooli", "applied")]
    r = client.post("/api/jobs/export", json={"tab": "applied"}, headers=H)
    assert len(list(load_workbook(io.BytesIO(r.content)).active.iter_rows())) == 3
    assert tr.stat().st_mtime_ns == before
    assert client.post("/api/jobs/export", json={"columns": ["nope"]}, headers=H).status_code == 400


def test_tracker_sync_and_open(client, opened, data):
    assert client.post("/api/tracker/sync", headers=H).json()["synced"] == 8
    r = client.post("/api/tracker/open", headers=H)
    assert r.status_code == 200 and opened == [["open", str(Tracker(settings=data["settings"]).path.resolve())]]


def test_open_refused_off_macos(client, data, monkeypatch):
    monkeypatch.setattr(desktop, "_platform", lambda: "linux")
    monkeypatch.setattr(desktop, "_run", lambda argv: pytest.fail("must not run"))
    r = client.post(f"/api/jobs/{data['jobs']['review']}/open-folder", headers=H)
    assert r.status_code == 409 and "macOS" in r.json()["detail"]


# --- Job detail --------------------------------------------------------------------------------------------------

def test_job_detail_extras_and_files(client, data, tmp_path):
    jid = data["jobs"]["review"]
    d = client.get(f"/api/jobs/{jid}").json()
    assert d["safety"]["flags"][0]["code"] == "GHOST_OLD_POST" and d["apply_session"]["outcome"] == "needs_review"
    assert d["screenshots"][0]["name"] == "01_form.png" and d["registry"] == {"verified": None, "flagged": None}
    assert d["activity"] and d["override"] is None
    r = client.get(f"/api/jobs/{jid}/files/resume.pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["x-content-type-options"] == "nosniff" and "inline" in r.headers["content-disposition"]
    assert client.get(f"/api/jobs/{jid}/files/cover_letter.md").headers["content-type"].startswith("text/plain")
    assert client.get(f"/api/jobs/{jid}/files/screenshots/01_form.png").headers["content-type"] == "image/png"
    outside = tmp_path / "secret.txt"
    outside.write_text("secret")
    os.symlink(outside, Store(data["settings"]).job_dir(jid) / "leak.txt")
    for bad in ("leak.txt", "..%2F..%2Fconfig%2Fpipeline.yaml", "../../config/pipeline.yaml", ".hidden", "nope.pdf"):
        assert client.get(f"/api/jobs/{jid}/files/{bad}").status_code == 404, bad
    assert client.get("/api/jobs/nope00000000/files/resume.pdf").status_code == 404


def test_status_writes_409_while_the_job_is_locked_and_applied_needs_submitted(client, data):
    from careeros.runs import locks

    s, jid = data["settings"], data["jobs"]["queued"]
    r = client.post(f"/api/jobs/{jid}/status", json={"status": "applied"}, headers=H)
    assert r.status_code == 400 and "Mark submitted" in r.json()["detail"]
    locks.acquire(RunStore(s).job_lock_path(jid), "run", 600, pid=os.getpid(), note="prepare")
    for path, body in (("status", {"status": "prepared"}), ("withdraw", {}), ("submitted", {})):
        r = client.post(f"/api/jobs/{jid}/{path}", json=body, headers=H)
        assert r.status_code == 409 and "locked" in r.json()["detail"], path
    assert Store(s).get_status(jid) == "queued"


def test_set_status_withdraw_undo_and_submitted(client, data):
    s, jid = data["settings"], data["jobs"]["queued"]
    r = client.post(f"/api/jobs/{jid}/status", json={"status": "prepared", "note": "by hand"}, headers=H)
    assert r.json() == {"status": "prepared", "previous": "queued"}
    assert Store(s).get_status(jid) == "prepared" and Tracker(settings=s).get_job(jid)["Status"] == "prepared"
    assert client.post(f"/api/jobs/{jid}/status", json={"status": "bogus"}, headers=H).status_code == 400
    r = client.post(f"/api/jobs/{jid}/withdraw", headers=H)
    assert r.json() == {"status": "withdrawn", "previous": "prepared"}
    client.post(f"/api/jobs/{jid}/status", json={"status": "prepared", "note": "undo withdraw"}, headers=H)
    assert Store(s).get_status(jid) == "prepared"
    assert client.post(f"/api/jobs/{jid}/submitted", headers=H).json()["status"] == "applied"
    assert Store(s)._read(jid, "status.json")["history"][-1]["note"].startswith("submitted by you")
    assert client.post("/api/jobs/nope00000000/status", json={"status": "queued"}, headers=H).status_code == 404


def test_override(client, data):
    s, jid = data["settings"], data["jobs"]["found"]
    got = client.post(f"/api/jobs/{jid}/override", json={"value": "manual"}, headers=H).json()
    assert got == {"override": "manual", "queued": False}
    assert Tracker(settings=s).read_overrides()[jid] == "manual"
    assert client.get(f"/api/jobs/{jid}").json()["override"] == "manual"
    assert client.post(f"/api/jobs/{jid}/override", json={"value": "X"}, headers=H).status_code == 400


def test_rerun_qa(client, data):
    r = client.post(f"/api/jobs/{data['jobs']['review']}/qa", headers=H)
    assert r.status_code == 200 and "pass" in r.json() and "summary" in r.json()


def test_safety_verify_flag_clear(client, data):
    s, jid = data["settings"], data["jobs"]["review"]
    r = client.post(f"/api/jobs/{jid}/safety/verify", json={"risk": "low", "signals": ["one"]}, headers=H)
    assert r.status_code == 400 and "two" in r.json()["detail"]
    r = client.post(f"/api/jobs/{jid}/safety/verify", json={"risk": "medium", "signals": ["careers page"],
                                                            "evidence": ["https://umbrella.example.com/careers"]},
                    headers=H)
    assert r.json()["risk"] == "medium"
    r = client.post(f"/api/jobs/{jid}/safety/flag", json={"reason": "manual", "confidence": "medium"}, headers=H)
    assert r.json()["state"] == "active"
    assert client.get(f"/api/jobs/{jid}").json()["registry"]["flagged"]["confidence"] == "medium"
    assert client.post(f"/api/jobs/{jid}/safety/clear", json={"note": "ok"}, headers=H).json()["state"] == "cleared"
    assert not registry.is_flagged(registry.load(registry.default_path(s)), "Umbrella Labs")
    assert client.post(f"/api/jobs/{data['jobs']['found']}/safety/clear", headers=H).status_code == 404
    assert client.post(f"/api/jobs/{jid}/safety/flag", json={"evidence": ["ftp://x"]}, headers=H).status_code == 400


def test_open_folder(client, data, opened):
    jid = data["jobs"]["review"]
    assert client.post(f"/api/jobs/{jid}/open-folder", headers=H).json() == {"opened": True}
    assert opened == [["open", str(Store(data["settings"]).job_dir(jid).resolve())]]
