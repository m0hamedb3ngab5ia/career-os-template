"""careeros.ui.services for the Today, Jobs and Job detail screens: open Action Items with done/undo, the jobs tabs
and xlsx export, Job detail extras (override, registry, activity), and every Job detail write going through the
existing domain code (Store / Tracker / safety registry / careeros.qa)."""
from __future__ import annotations

import io
import os
import sys
from datetime import datetime, timezone

import pytest
from conftest import make_temp_root
from fixtures.ui_data import build_ui_data
from openpyxl import load_workbook

from careeros.safety import registry
from careeros.store import Store
from careeros.tracker import Tracker
from careeros.ui.config import load_ui_config
from careeros.ui.index import Index
from careeros.ui.services import desktop
from careeros.ui.services import job_actions as acts
from careeros.ui.services import jobs as jobs_svc
from careeros.ui.services import status as status_svc
from careeros.ui.services import today as today_svc

pytestmark = pytest.mark.unit

NOW = datetime(2026, 9, 24, 15, 0, tzinfo=timezone.utc)


@pytest.fixture
def data(tmp_path):
    return build_ui_data(make_temp_root(tmp_path / "repo"), NOW)


@pytest.fixture
def idx(data):
    ix = Index(data["settings"])
    ix.rebuild()
    yield ix
    ix.close()


@pytest.fixture
def opened(monkeypatch):
    calls: list[list[str]] = []
    monkeypatch.setattr(desktop, "_run", lambda argv: calls.append(list(argv)))
    monkeypatch.setattr(desktop, "_platform", lambda: "darwin")
    return calls


def closed(data):
    return load_ui_config(data["settings"]).closed


# --- Today -------------------------------------------------------------------------------------------------------

def test_open_actions_lists_every_open_item_with_empty_due(data, idx):
    items = today_svc.open_actions(idx)
    assert [i["company"] for i in items] == ["Umbrella Labs", "Initech", "Stark Industries"]
    assert all(i["due"] is None and i["due_reason"] is None for i in items)
    assert items[0]["link"] == "https://boards.example.com/review" and items[0]["needs"] == "laptop"


def test_mark_done_and_reopen_go_through_the_tracker(data):
    s, aid = data["settings"], data["actions"]["high"]
    assert today_svc.mark_done(s, aid) == {"ok": True, "queued": False}
    assert aid not in {i["ID"] for i in Tracker(settings=s).list_action_items()}
    assert today_svc.reopen(s, aid) == {"ok": True, "queued": False}
    assert aid in {i["ID"] for i in Tracker(settings=s).list_action_items()}
    with pytest.raises(LookupError):
        today_svc.mark_done(s, "nope")
    with pytest.raises(LookupError):
        today_svc.reopen(s, "nope")


def test_prepare_queue_counts_real_candidates(data):
    q = today_svc.prepare_queue(data["settings"], NOW)
    assert q["error"] is None and isinstance(q["total"], int) and q["total"] >= 0


def test_response_rate_breakdown(data, idx):
    rr = status_svc.status(data["settings"], idx, NOW)["tiles"]["response_rate"]
    by = {b["status"]: b for b in rr["breakdown"]}
    assert by["interview"]["count"] == 1 and by["interview"]["companies"] == ["Stark Industries"]
    assert by["rejected"]["count"] == 1 and by["no_reply"]["count"] == 1
    assert by["no_reply"]["companies"] == ["Hooli"]
    assert sum(b["count"] for b in rr["breakdown"]) == rr["applied"]


# --- Jobs list ---------------------------------------------------------------------------------------------------

def test_tabs_count_each_view(data, idx):
    tabs = {t["key"]: t for t in jobs_svc.tabs(idx, closed(data))}
    assert list(tabs) == ["active", "review", "applied", "tier_a", "all"]
    assert tabs["all"]["count"] == 8 and tabs["review"]["count"] == 1
    assert tabs["applied"]["count"] == 2            # applied + interview
    assert tabs["active"]["count"] == 6             # not rejected / skipped
    assert tabs["tier_a"]["count"] == 2 and tabs["tier_a"]["label"] == "Tier A"
    assert {t["key"]: t["count"] for t in jobs_svc.tabs(idx, closed(data), q="hooli")}["all"] == 1


def test_list_by_tab_and_next_action(data, idx):
    page = jobs_svc.list_jobs(idx, tab="applied", closed=closed(data))
    assert {j["company"] for j in page["items"]} == {"Hooli", "Stark Industries"}
    rows = {j["company"]: j for j in jobs_svc.list_jobs(idx)["items"]}
    assert rows["Umbrella Labs"]["next_action"] == "Review and submit"
    assert rows["Hooli"]["next_action"] is None          # its only item is done
    with pytest.raises(ValueError):
        jobs_svc.list_jobs(idx, tab="bogus", closed=closed(data))


def test_export_xlsx_selected_rows_and_columns(data, idx):
    ids = [data["jobs"]["review"], data["jobs"]["applied"]]
    raw = jobs_svc.export_xlsx(idx, job_ids=ids, columns=["company", "fit", "status"])
    ws = load_workbook(io.BytesIO(raw)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("JobID", "Company", "Fit", "Status")
    assert {r[1] for r in rows[1:]} == {"Umbrella Labs", "Hooli"}
    assert not Store(data["settings"]).job_dir(ids[0]).joinpath("export.xlsx").exists()


def test_export_xlsx_by_filter_and_bad_input(data, idx):
    raw = jobs_svc.export_xlsx(idx, tab="applied", closed=closed(data))
    rows = list(load_workbook(io.BytesIO(raw)).active.iter_rows(values_only=True))
    assert len(rows) == 3 and "Next action" in rows[0]
    with pytest.raises(ValueError):
        jobs_svc.export_xlsx(idx, columns=["bogus"])
    with pytest.raises(ValueError):
        jobs_svc.export_xlsx(idx, job_ids=["../etc"])


def test_export_neutralises_formula_text(data, idx):
    idx.con.execute("UPDATE jobs SET title = '=HYPERLINK(\"http://x\")' WHERE job_id = ?", (data["jobs"]["found"],))
    raw = jobs_svc.export_xlsx(idx, job_ids=[data["jobs"]["found"]], columns=["title"])
    ws = load_workbook(io.BytesIO(raw)).active
    assert ws.cell(row=2, column=2).value.startswith("'=")


# --- Job detail extras -------------------------------------------------------------------------------------------

def test_job_detail_extras(data, idx):
    s = data["settings"]
    jid = data["jobs"]["interview"]
    Tracker(settings=s).upsert_job({"job_id": jid, "override": "manual"})
    registry.add_verified(registry.verified_path(s), "Stark Industries", "low", signals=["site", "ats"])
    d = jobs_svc.job_detail(s, idx, jid)
    assert d["override"] == "manual"
    assert d["registry"]["verified"]["risk"] == "low" and d["registry"]["flagged"] is None
    assert [c["manual"] for c in d["contacts_policy"]] == [True, False]
    assert d["activity"][0]["component"] == "fixture" and "status -> interview" in d["activity"][0]["message"]
    assert d["outreach"] is None


def test_job_detail_override_none_without_tracker_row(data, idx):
    d = jobs_svc.job_detail(data["settings"], idx, data["jobs"]["found"])
    assert d["override"] is None


# --- files -------------------------------------------------------------------------------------------------------

def test_resolve_file_inside_the_job_dir_only(data, tmp_path):
    s, jid = data["settings"], data["jobs"]["review"]
    jd = Store(s).job_dir(jid)
    assert acts.resolve_file(s, jid, "resume.pdf") == (jd / "resume.pdf").resolve()
    assert acts.resolve_file(s, jid, "screenshots/01_form.png").name == "01_form.png"
    secret = tmp_path / "secret.txt"
    secret.write_text("no")
    os.symlink(secret, jd / "link.txt")
    os.symlink(tmp_path, jd / "linkdir")
    for bad in ("../../config/pipeline.yaml", "/etc/passwd", "link.txt", "linkdir/secret.txt", ".hidden",
                "screenshots/../../x", "missing.pdf", "screenshots", ""):
        with pytest.raises(LookupError):
            acts.resolve_file(s, jid, bad)
    with pytest.raises(LookupError):
        acts.resolve_file(s, "../x", "resume.pdf")


def test_resolve_file_refuses_a_symlink_into_a_hidden_path(data):
    s, jid = data["settings"], data["jobs"]["review"]
    jd = Store(s).job_dir(jid)
    (jd / ".private").mkdir()
    (jd / ".private" / "notes.txt").write_text("no")
    (jd / ".env").write_text("no")
    os.symlink(jd / ".private" / "notes.txt", jd / "notes.txt")
    os.symlink(jd / ".env", jd / "env.txt")
    os.symlink(jd / ".private", jd / "priv")
    for bad in ("notes.txt", "env.txt", "priv/notes.txt"):
        with pytest.raises(LookupError):
            acts.resolve_file(s, jid, bad)


# --- writes ------------------------------------------------------------------------------------------------------

def _tracker_status(s, jid):
    return (Tracker(settings=s).get_job(jid) or {}).get("Status")


def test_set_status_writes_status_json_and_tracker(data):
    s, jid = data["settings"], data["jobs"]["queued"]
    out = acts.set_status(s, jid, "needs_review", "checked by hand")
    assert out == {"status": "needs_review", "previous": "queued"}
    assert Store(s).get_status(jid) == "needs_review" and _tracker_status(s, jid) == "needs_review"
    with pytest.raises(ValueError):
        acts.set_status(s, jid, "bogus")
    with pytest.raises(LookupError):
        acts.set_status(s, "nope00000000", "queued")


def _hold_lock(s, jid):
    from careeros.runs import locks
    from careeros.runs.store import RunStore

    return locks.acquire(RunStore(s).job_lock_path(jid), "run", 600, pid=os.getpid(), note="prepare")


def test_status_writes_refuse_while_the_job_is_locked(data):
    s, jid = data["settings"], data["jobs"]["queued"]
    _hold_lock(s, jid)
    for write in (lambda: acts.set_status(s, jid, "needs_review"), lambda: acts.withdraw(s, jid),
                  lambda: acts.mark_submitted(s, jid)):
        with pytest.raises(acts.JobLocked):
            write()
    assert Store(s).get_status(jid) == "queued"


def test_set_status_refuses_applied_except_to_undo_a_withdraw(data):
    s, jid = data["settings"], data["jobs"]["queued"]
    with pytest.raises(ValueError, match="Mark submitted"):
        acts.set_status(s, jid, "applied")
    acts.withdraw(s, jid)
    with pytest.raises(ValueError):
        acts.set_status(s, jid, "applied")  # was queued before the withdraw: not an undo
    assert Store(s).get_status(jid) == "withdrawn"


def test_withdraw_then_undo(data):
    s, jid = data["settings"], data["jobs"]["applied"]
    assert acts.withdraw(s, jid) == {"status": "withdrawn", "previous": "applied"}
    assert Store(s).get_status(jid) == "withdrawn"
    acts.set_status(s, jid, "applied", "undo withdraw")
    assert Store(s).get_status(jid) == "applied"


def test_set_status_undo_restores_applied_only_right_after_leaving_it(data):
    s, jid = data["settings"], data["jobs"]["applied"]
    assert acts.set_status(s, jid, "interview") == {"status": "interview", "previous": "applied"}
    acts.set_status(s, jid, "applied", "undo status change")  # the Undo toast of the change just made
    assert Store(s).get_status(jid) == "applied"
    acts.set_status(s, jid, "interview")
    acts.set_status(s, jid, "offer")
    with pytest.raises(ValueError, match="Mark submitted"):
        acts.set_status(s, jid, "applied")  # the latest change was interview -> offer: not an undo
    q = data["jobs"]["queued"]
    acts.set_status(s, q, "needs_review")
    with pytest.raises(ValueError, match="Mark submitted"):
        acts.set_status(s, q, "applied")  # never applied
    assert Store(s).get_status(q) == "needs_review"


def test_set_override_reports_queued_while_excel_holds_the_tracker(data, monkeypatch):
    from openpyxl.workbook.workbook import Workbook

    s, jid = data["settings"], data["jobs"]["queued"]
    acts.set_override(s, jid, "B")  # the Jobs row exists

    def locked(self, filename):
        raise PermissionError(13, "locked")

    monkeypatch.setattr(Workbook, "save", locked)
    with pytest.warns(UserWarning):
        assert acts.set_override(s, jid, "skip") == {"override": "skip", "queued": True}


def test_job_dir_for_refuses_a_symlinked_folder_outside_jobs_dir(data, tmp_path):
    s, jid = data["settings"], data["jobs"]["queued"]
    outside = tmp_path / "outside" / "evil01"
    outside.mkdir(parents=True)
    (outside / "posting.json").write_text((Store(s).job_dir(jid) / "posting.json").read_text())
    os.symlink(outside, Store(s).job_dir(jid).parent / "evil01")
    assert jobs_svc.job_dir_for(s, "evil01") is None
    assert jobs_svc.job_dir_for(s, jid) is not None


def test_mark_submitted_sets_applied_with_note(data):
    s, jid = data["settings"], data["jobs"]["review"]
    assert acts.mark_submitted(s, jid)["status"] == "applied"
    st = Store(s)._read(jid, "status.json")
    assert st["status"] == "applied" and "submitted" in st["history"][-1]["note"]


def test_set_override_writes_the_tracker_column(data):
    s, jid = data["settings"], data["jobs"]["queued"]
    assert acts.set_override(s, jid, "skip") == {"override": "skip", "queued": False}
    assert Tracker(settings=s).read_overrides()[jid] == "skip"
    acts.set_override(s, jid, "")
    assert jid not in Tracker(settings=s).read_overrides()
    with pytest.raises(ValueError):
        acts.set_override(s, jid, "Z")


def test_rerun_qa_runs_the_deterministic_checks_and_logs(data):
    s, jid = data["settings"], data["jobs"]["review"]
    out = acts.rerun_qa(s, jid)
    assert "pass" in out and "summary" in out and "checks" in out
    assert "qa (deterministic)" in Store(s).read_log(jid)


def test_safety_verify_flag_clear_use_the_registry(data):
    s, jid = data["settings"], data["jobs"]["review"]
    e = acts.safety_verify(s, jid, risk="low", signals=["official site", "ATS listing"],
                           evidence=["https://umbrella.example.com"])
    assert e["company"] == "Umbrella Labs" and e["risk"] == "low"
    with pytest.raises(ValueError):
        acts.safety_verify(s, jid, risk="bogus", signals=[], evidence=[])
    f = acts.safety_flag(s, jid, reason="SCAM_PAYMENT_REQUEST", confidence="medium",
                         evidence=["https://boards.example.com/review"])
    assert f["confidence"] == "medium" and f["state"] == "active"
    assert registry.is_flagged(registry.load(registry.default_path(s)), "Umbrella Labs")
    c = acts.safety_clear(s, jid, note="looked fine")
    assert c["state"] == "cleared"
    with pytest.raises(ValueError):
        acts.safety_flag(s, jid, confidence="low")
    with pytest.raises(ValueError):
        acts.safety_flag(s, jid, evidence=["javascript:alert(1)"])
    with pytest.raises(LookupError):
        acts.safety_clear(s, data["jobs"]["found"])


def test_open_folder_and_tracker_only_their_own_paths(data, opened):
    s, jid = data["settings"], data["jobs"]["review"]
    acts.open_folder(s, jid)
    assert opened == [["open", str(Store(s).job_dir(jid).resolve())]]
    with pytest.raises(LookupError):
        acts.open_folder(s, "../x")
    out = acts.open_tracker(s)
    assert opened[-1] == ["open", str(Tracker(settings=s).path.resolve())] and out["opened"] is True


def test_open_refuses_off_macos(data, monkeypatch):
    monkeypatch.setattr(desktop, "_platform", lambda: "linux")
    with pytest.raises(desktop.Unsupported):
        acts.open_folder(data["settings"], data["jobs"]["review"])


def test_open_tracker_missing_file(data, opened):
    Tracker(settings=data["settings"]).path.unlink()
    with pytest.raises(desktop.Unsupported):
        acts.open_tracker(data["settings"])
    assert opened == []


def test_sync_tracker(data):
    out = acts.sync_tracker(data["settings"])
    assert out["synced"] == 8 and out["pending"] == 0


@pytest.mark.skipif(sys.platform == "win32", reason="posix paths")
def test_platform_default_is_sys_platform():
    assert desktop._platform() == sys.platform
