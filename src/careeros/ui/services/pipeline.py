"""GET /api/pipeline: the board. One column per `pipeline.yaml: ui.pipeline.columns` entry with its count and its
first `card_limit` cards (all of them for the columns named in `expand`), the Closed line (statuses in no column),
and the filter options. Cards come from the index; the override comes from the tracker's Jobs tab (read-only) and
the hint line from the job's open Action Item, else its safety flag, else its stage.
"""
from __future__ import annotations

import json
import threading
from pathlib import Path
from typing import Any, Literal, Union

from typing_extensions import TypedDict

from careeros.ui.config import load_ui_config

CARD_FIELDS = ("job_id", "company", "title", "location", "category", "fit", "tier", "status", "safety", "qa_passed",
               "qa_score", "found_at", "updated_at")
FILTERS = ("tier", "category", "safety", "location")
_PRIO = {"H": 0, "M": 1, "L": 2}
PREPARE_STAGE = ("queued", "prepared", "needs_review")


# Response shapes (OpenAPI -> ui/src/api/schema.gen.ts). Index columns are nullable TEXT/INTEGER.
class ActionHint(TypedDict):
    kind: Literal["action"]
    type: str
    due: str | None
    due_reason: str | None


class SafetyHint(TypedDict):
    kind: Literal["safety"]
    text: str


class NotScoredHint(TypedDict):
    kind: Literal["not_scored"]


class TierAHint(TypedDict):
    kind: Literal["tier_a"]


class QaFailedHint(TypedDict):
    kind: Literal["qa_failed"]


Hint = Union[ActionHint, SafetyHint, NotScoredHint, TierAHint, QaFailedHint]


class Card(TypedDict):
    job_id: str
    company: str | None
    title: str | None
    location: str | None
    category: str | None
    fit: int | None
    tier: str | None
    status: str
    safety: str | None
    qa_passed: bool | None
    qa_score: float | None
    found_at: str | None
    updated_at: str | None
    override: str | None
    hint: Hint | None


class Column(TypedDict):
    name: str
    statuses: list[str]
    count: int
    cards: list[Card]


class Closed(TypedDict):
    count: int
    by_status: dict[str, int]


class LocationOption(TypedDict):
    value: str
    count: int


class BoardOptions(TypedDict):
    categories: list[str]
    locations: list[LocationOption]


class Board(TypedDict):
    columns: list[Column]
    closed: Closed
    card_limit: int
    options: BoardOptions


def _where(filters: dict[str, list[str] | None]) -> tuple[str, list[Any]]:
    parts, params = [], []
    for col in FILTERS:
        vals = [v for v in filters.get(col) or [] if v is not None]
        if vals:
            parts.append(f"{col} IN ({','.join('?' * len(vals))})")
            params.extend(vals)
    return (" AND ".join(parts), params)


# --- tracker overrides (read-only, cached on the file's signature) ------------------------------------------------

_cache_lock = threading.Lock()
_override_cache: dict[str, tuple[str, dict[str, str]]] = {}


def read_overrides(tracker_path: Path) -> dict[str, str]:
    """{job_id: Override} from the Jobs tab, opened read-only (never through Tracker, which may create or repair
    the workbook). A missing or unreadable workbook means no overrides."""
    p = Path(tracker_path)
    try:
        st = p.stat()
    except OSError:
        return {}
    sig = f"{st.st_mtime_ns}:{st.st_size}"
    with _cache_lock:
        hit = _override_cache.get(str(p))
        if hit and hit[0] == sig:
            return hit[1]
    out: dict[str, str] = {}
    try:
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            rows = wb["Jobs"].iter_rows(values_only=True)
            header = [str(h) if h is not None else "" for h in next(rows, ())]
            if "JobID" in header and "Override" in header:
                ji, oi = header.index("JobID"), header.index("Override")
                for r in rows:
                    if r and len(r) > max(ji, oi) and r[ji] not in (None, "") and r[oi] not in (None, ""):
                        out[str(r[ji])] = str(r[oi]).strip()
        finally:
            wb.close()
    except Exception:  # noqa: BLE001 - locked, half-written or foreign workbook: no overrides this time
        return {}
    with _cache_lock:
        _override_cache[str(p)] = (sig, out)
    return out


# --- hints --------------------------------------------------------------------------------------------------------

def _safety_text(jobs_dir: Path, job_id: str) -> str | None:
    try:
        data = json.loads((jobs_dir / job_id / "safety.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    flags = [f for f in data.get("flags") or [] if isinstance(f, dict) and f.get("level") != "info"]
    order = {"block": 0, "skip": 1, "review": 2}
    flags.sort(key=lambda f: order.get(str(f.get("level")), 3))
    for f in flags:
        text = str(f.get("detail") or "").strip()
        if text:
            return text[:80]
    return None


def card_hint(card: dict[str, Any], action: dict[str, Any] | None, safety_text: str | None) -> Hint | None:
    """The one-line "what's next" under a card: an open Action Item, else a safety flag, else the stage."""
    if action:
        return {"kind": "action", "type": action.get("type") or "other", "due": action.get("due"),
                "due_reason": action.get("due_reason")}
    if safety_text and card.get("safety") in ("block", "review", "skip"):
        return {"kind": "safety", "text": safety_text}
    if card.get("status") == "found" and card.get("fit") is None:
        return {"kind": "not_scored"}
    if card.get("tier") == "A" and card.get("status") in PREPARE_STAGE:
        return {"kind": "tier_a"}
    if card.get("qa_passed") is False:
        return {"kind": "qa_failed"}
    return None


def _open_actions(ix: Any, job_ids: list[str]) -> dict[str, dict[str, Any]]:
    """The most pressing open item per job: highest priority, then soonest due, then oldest."""
    if not job_ids:
        return {}
    best: dict[str, dict[str, Any]] = {}
    rows = ix.query(f"SELECT job_id, type, priority, due, due_reason, created FROM action_items "
                    f"WHERE done = 0 AND job_id IN ({','.join('?' * len(job_ids))})", job_ids)

    def key(r: dict[str, Any]) -> tuple[Any, ...]:
        return (_PRIO.get(r.get("priority") or "", 3), r.get("due") is None, r.get("due") or "", r.get("created") or "")

    for r in rows:
        cur = best.get(r["job_id"])
        if cur is None or key(r) < key(cur):
            best[r["job_id"]] = r
    return best


# --- the board ----------------------------------------------------------------------------------------------------

def board(settings: Any, ix: Any, *, tier: list[str] | None = None, category: list[str] | None = None,
          safety: list[str] | None = None, location: list[str] | None = None,
          expand: list[str] | None = None) -> Board:
    ui = load_ui_config(settings)
    where, params = _where({"tier": tier, "category": category, "safety": safety, "location": location})
    clause = f" AND {where}" if where else ""
    by = {r["status"]: r["n"] for r in ix.query(f"SELECT status, COUNT(*) AS n FROM jobs WHERE 1=1{clause} "
                                                "GROUP BY status", params)}
    expand_set = set(expand or [])
    overrides = read_overrides(Path(settings.paths["tracker_xlsx"]))
    jobs_dir = Path(settings.paths["jobs_dir"])
    columns = []
    for c in ui.columns:
        sts = c["statuses"]
        limit = "" if c["name"] in expand_set else f" LIMIT {int(ui.card_limit)}"
        rows = ix.query(f"SELECT {', '.join(CARD_FIELDS)} FROM jobs WHERE status IN ({','.join('?' * len(sts))})"
                        f"{clause} ORDER BY (fit IS NULL), fit DESC, updated_at DESC, company ASC, job_id ASC{limit}",
                        [*sts, *params])
        columns.append({"name": c["name"], "statuses": list(sts), "count": sum(by.get(s, 0) for s in sts),
                        "cards": rows})
    acts = _open_actions(ix, [r["job_id"] for col in columns for r in col["cards"]])
    for col in columns:
        cards = []
        for r in col["cards"]:
            card = {**r, "qa_passed": None if r["qa_passed"] is None else bool(r["qa_passed"]),
                    "override": overrides.get(r["job_id"]) or None}
            action = acts.get(r["job_id"])
            text = _safety_text(jobs_dir, r["job_id"]) if not action and r["safety"] in ("block", "review", "skip") \
                else None
            card["hint"] = card_hint(card, action, text)
            cards.append(card)
        col["cards"] = cards
    closed = {s: by[s] for s in ui.closed if by.get(s)}
    return {
        "columns": columns,
        "closed": {"count": sum(closed.values()), "by_status": closed},
        "card_limit": ui.card_limit,
        "options": {
            "categories": [r["category"] for r in ix.query(
                "SELECT DISTINCT category FROM jobs WHERE category IS NOT NULL AND category != '' ORDER BY category")],
            "locations": [{"value": r["location"], "count": r["n"]} for r in ix.query(
                "SELECT location, COUNT(*) AS n FROM jobs WHERE location IS NOT NULL AND location != '' "
                "GROUP BY location ORDER BY n DESC, location ASC")],
        },
    }
