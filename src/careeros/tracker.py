from __future__ import annotations

import json
import os
import re
import uuid
import warnings
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Callable, Iterator
from zipfile import BadZipFile

try:
    import fcntl
except ImportError:  # Windows: no cross-process lock
    fcntl = None  # type: ignore[assignment]

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from careeros.config import Settings, get_settings, normalize_company
from careeros.models import ACTION_NEEDS, ACTION_TYPES, STATUSES, ActionItem, Contact, TrackerRow

MAX_ROWS = 5000

JOB_COLUMNS: list[tuple[str, str, int]] = [
    ("JobID", "job_id", 14),
    ("Company", "company", 22),
    ("Role", "role", 40),
    ("Category", "category", 16),
    ("Tier", "tier", 6),
    ("Prestige", "prestige", 10),
    ("Fit", "fit", 6),
    ("Location", "location", 24),
    ("Remote", "remote", 8),
    ("Salary", "salary", 18),
    ("URL", "url", 40),
    ("DateFound", "date_found", 12),
    ("DateApplied", "date_applied", 12),
    ("Status", "status", 14),
    ("ATS", "ats", 11),
    ("ResumeVersion", "resume_version", 14),
    ("CoverLetter", "cover_letter", 11),
    ("QAScore", "qa_score", 9),
    ("Override", "override", 10),
    ("LastEmailDate", "last_email_date", 13),
    ("NextAction", "next_action", 28),
    ("NextActionDate", "next_action_date", 14),
    ("Notes", "notes", 40),
    ("Folder", "folder", 12),
]
ACTION_COLUMNS = [
    ("ID", 10), ("Created", 18), ("JobID", 14), ("Company", 22), ("Role", 36), ("Type", 14),
    ("What to do", 50), ("Link", 40), ("Priority", 9), ("Needs", 9), ("Done", 7), ("DoneDate", 12),
]
CONTACT_COLUMNS = [
    ("JobID", 14), ("Company", 22), ("Name", 24), ("Title", 28), ("LinkedIn", 40), ("Email", 30),
    ("EmailConfidence", 15), ("DraftMessage", 60), ("Sent", 7), ("SentDate", 12), ("Replied", 10),
]
LOG_COLUMNS = [("Timestamp", 20), ("JobID", 14), ("Component", 14), ("Message", 90)]
CONFIG_COLUMNS = [("Key", 24), ("Value", 60)]

HEADER_FILL = PatternFill("solid", fgColor="DDE5F0")

# `careeros tracker upsert <job_id> --field k=v`: accept the snake_case key or the column header.
_JOB_KEYS = {key for _, key, _ in JOB_COLUMNS}
_JOB_KEY_BY_HEADER = {header.lower(): key for header, key, _ in JOB_COLUMNS}
_DATE_KEYS = {"date_found", "date_applied", "last_email_date", "next_action_date"}


def parse_field_args(pairs: list[str]) -> dict[str, Any]:
    """["DateApplied=today", "status=applied", "Fit=88"] -> {"date_applied": "YYYY-MM-DD", ...}.

    Keys: Jobs column header (case-insensitive) or its snake_case name. `today` in a date column
    becomes today's date. Status must be a lifecycle status; Fit an int; QAScore a float.
    Raises ValueError with a readable message on anything else.
    """
    out: dict[str, Any] = {}
    for pair in pairs:
        if "=" not in pair:
            raise ValueError(f"--field expects key=value, got {pair!r}")
        raw_key, value = (x.strip() for x in pair.split("=", 1))
        key = raw_key if raw_key in _JOB_KEYS else _JOB_KEY_BY_HEADER.get(raw_key.lower())
        if key is None:
            valid = ", ".join(h for h, _, _ in JOB_COLUMNS if h != "JobID")
            raise ValueError(f"unknown field {raw_key!r}; valid: {valid}")
        if key == "job_id":
            raise ValueError("JobID is the positional argument, not a --field")
        val: Any = value
        if key in _DATE_KEYS and value.lower() == "today":
            val = _today()
        elif key == "status" and value not in STATUSES:
            raise ValueError(f"invalid status {value!r}; expected one of {', '.join(STATUSES)}")
        elif key == "fit":
            try:
                val = int(value)
            except ValueError:
                raise ValueError(f"Fit must be an integer, got {value!r}") from None
        elif key == "qa_score":
            try:
                val = float(value)
            except ValueError:
                raise ValueError(f"QAScore must be a number, got {value!r}") from None
        out[key] = val
    return out


def _style_header(ws: Worksheet, columns: list[tuple[str, int]]) -> None:
    for i, (name, width) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{MAX_ROWS}"


def _add_list_validation(ws: Worksheet, col_idx: int, values: list[str]) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error = "Pick a value from the list"
    dv.errorTitle = "Invalid"
    col = get_column_letter(col_idx)
    dv.add(f"{col}2:{col}{MAX_ROWS}")
    ws.add_data_validation(dv)


def _refresh_type_validation(ws: Worksheet, hdr: dict[str, int]) -> None:
    """Keep the Action Items 'Type' dropdown in sync with ACTION_TYPES (new types are appended over time)."""
    if "Type" not in hdr:
        return
    col = get_column_letter(hdr["Type"])
    want = '"' + ",".join(ACTION_TYPES) + '"'
    dvs = ws.data_validations.dataValidation
    for dv in list(dvs):
        if dv.type == "list" and any(str(r).startswith(f"{col}2") for r in str(dv.sqref).split()):
            if dv.formula1 == want:
                return
            dvs.remove(dv)
    _add_list_validation(ws, hdr["Type"], list(ACTION_TYPES))


def _sanitize_urls(ws: Worksheet) -> None:
    """URL column: text only (never a formula); hyperlinks only for http(s)."""
    col = _header_index(ws).get("URL")
    if col is None:
        return
    for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        if cell.data_type == "f":
            cell.data_type = "s"
        link = cell.hyperlink.target if cell.hyperlink is not None else None
        if link is not None and not re.match(r"https?://", str(link), re.I):
            cell.hyperlink = None


def _header_index(ws: Worksheet) -> dict[str, int]:
    return {str(c.value): c.column for c in ws[1] if c.value is not None}


def _find_row(ws: Worksheet, col_idx: int, value: str) -> int | None:
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        cell = row[0]
        if cell.value is not None and str(cell.value) == value:
            return cell.row
    return None


def _put(ws: Worksheet, row: int, column: int, value: Any) -> Any:
    """Write a cell. Strings are always text: board data like "=HYPERLINK(...)" must never become a formula."""
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def _today() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


class Tracker:
    def __init__(self, path: Path | None = None, settings: Settings | None = None):
        s = settings if settings is not None else (None if path else get_settings())
        self.settings = s
        self.path = Path(path) if path else s.paths["tracker_xlsx"]  # type: ignore[union-attr]
        self.pending_path = self.path.with_name(self.path.name + ".pending.json")
        self.jobs_dir: Path | None = s.paths["jobs_dir"] if s else None
        self.lock_path = self.path.with_name(f".{self.path.name}.lock")
        self._lock_depth = 0

    @contextmanager
    def _lock(self) -> Iterator[None]:
        """Cross-process exclusive lock around load -> mutate -> save and pending-queue writes.
        Reentrant within one Tracker (flush replays ops that lock again)."""
        if self._lock_depth or fcntl is None:
            self._lock_depth += 1
            try:
                yield
            finally:
                self._lock_depth -= 1
            return
        self.lock_path.parent.mkdir(parents=True, exist_ok=True)
        with self.lock_path.open("a") as fh:
            fcntl.flock(fh, fcntl.LOCK_EX)
            self._lock_depth += 1
            try:
                yield
            finally:
                self._lock_depth -= 1
                fcntl.flock(fh, fcntl.LOCK_UN)

    # --- create ------------------------------------------------------------

    def init(self, force: bool = False) -> Path:
        if self.path.exists() and not force:
            return self.path
        wb = Workbook()
        ws = wb.active
        ws.title = "Jobs"
        _style_header(ws, [(n, w) for n, _, w in JOB_COLUMNS])
        col = {n: i for i, (n, _, _) in enumerate(JOB_COLUMNS, start=1)}
        _add_list_validation(ws, col["Status"], list(STATUSES))
        _add_list_validation(ws, col["Tier"], ["A", "B", "C"])
        _add_list_validation(ws, col["Remote"], ["Y", "N"])
        _add_list_validation(ws, col["CoverLetter"], ["Y", "N"])
        _add_list_validation(ws, col["Override"], ["", "A", "B", "C", "skip", "manual"])

        wa = wb.create_sheet("Action Items")
        _style_header(wa, ACTION_COLUMNS)
        acol = {n: i for i, (n, _) in enumerate(ACTION_COLUMNS, start=1)}
        _add_list_validation(wa, acol["Type"], list(ACTION_TYPES))
        _add_list_validation(wa, acol["Priority"], ["H", "M", "L"])
        _add_list_validation(wa, acol["Needs"], list(ACTION_NEEDS))
        _add_list_validation(wa, acol["Done"], ["Y", "N"])

        wc = wb.create_sheet("Contacts")
        _style_header(wc, CONTACT_COLUMNS)
        ccol = {n: i for i, (n, _) in enumerate(CONTACT_COLUMNS, start=1)}
        _add_list_validation(wc, ccol["Sent"], ["Y", "N"])

        wl = wb.create_sheet("Log")
        _style_header(wl, LOG_COLUMNS)

        wk = wb.create_sheet("Config")
        _style_header(wk, CONFIG_COLUMNS)
        for i, (k, v) in enumerate(
            [("created", _now()), ("last_scout", ""), ("last_sync", ""), ("jobs_count", 0), ("applied_count", 0)],
            start=2,
        ):
            wk.cell(row=i, column=1, value=k)
            wk.cell(row=i, column=2, value=v)

        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._save(wb)
        return self.path

    # --- core mutate / read -----------------------------------------------

    def _load(self) -> Workbook:
        if not self.path.exists():
            self.init()
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = load_workbook(self.path)
            self._migrate(wb)
            return wb
        except (BadZipFile, KeyError, InvalidFileException) as e:
            # Only a malformed file is "corrupt". OSErrors (locks, iCloud) propagate so _mutate queues them.
            backup = self.path.with_name(f"{self.path.stem}.corrupt-{datetime.now():%Y%m%d-%H%M%S}{self.path.suffix}")
            self.path.replace(backup)
            warnings.warn(f"tracker {self.path.name} unreadable ({e}); moved to {backup.name} and recreated", stacklevel=3)
            self.init()
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                return load_workbook(self.path)

    @staticmethod
    def _migrate(wb: Workbook) -> None:
        """Add columns introduced after a workbook was created (non-destructive: appended after existing headers),
        and neutralise Jobs URL cells written before they were text-safe."""
        if "Jobs" in wb.sheetnames:
            _sanitize_urls(wb["Jobs"])
        if "Action Items" not in wb.sheetnames:
            return
        ws = wb["Action Items"]
        hdr = _header_index(ws)
        _refresh_type_validation(ws, hdr)
        if "Needs" in hdr:
            return
        # never move existing columns: append at the end so user data stays put
        col = max(hdr.values(), default=0) + 1
        c = ws.cell(row=1, column=col, value="Needs")
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 9
        _add_list_validation(ws, col, list(ACTION_NEEDS))
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=hdr["ID"]).value not in (None, ""):
                _put(ws, r, col, "anytime")

    def _save(self, wb: Workbook) -> None:
        tmp = self.path.with_name(f".{self.path.stem}.tmp-{os.getpid()}{self.path.suffix}")
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        finally:
            if tmp.exists():
                tmp.unlink(missing_ok=True)

    def _mutate(self, op: str, payload: dict[str, Any], fn: Callable[[Workbook], Any]) -> Any:
        with self._lock():
            return self._mutate_locked(op, payload, fn)

    def _mutate_locked(self, op: str, payload: dict[str, Any], fn: Callable[[Workbook], Any]) -> Any:
        try:
            wb = self._load()
            result = fn(wb)
            self._save(wb)
            return result
        except PermissionError as e:
            self._queue(op, payload, str(e))
            return None
        except OSError as e:
            if e.errno in (13, 16, 26, 35):
                self._queue(op, payload, str(e))
                return None
            raise

    def _queue(self, op: str, payload: dict[str, Any], err: str) -> None:
        q = self._read_pending()
        q.append({"op": op, "payload": payload, "queued_at": _now()})
        self._write_pending(q)
        warnings.warn(
            f"tracker locked ({err}); queued '{op}' to {self.pending_path.name} — run `careeros tracker flush`",
            stacklevel=3,
        )

    def _read_pending(self) -> list[dict[str, Any]]:
        if not self.pending_path.exists():
            return []
        try:
            return json.loads(self.pending_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return []

    def _write_pending(self, q: list[dict[str, Any]]) -> None:
        if not q:
            self.pending_path.unlink(missing_ok=True)
            return
        tmp = self.pending_path.with_name(self.pending_path.name + ".tmp")
        tmp.write_text(json.dumps(q, indent=2, default=str), encoding="utf-8")
        tmp.replace(self.pending_path)

    def pending_count(self) -> int:
        return len(self._read_pending())

    def flush_pending(self) -> int:
        """Replay queued ops in order. If the workbook is still locked (the op re-queues itself) or an op
        raises, the failed op and every op after it stay queued, in their original order."""
        with self._lock():
            return self._flush_locked()

    def _flush_locked(self) -> int:
        q = self._read_pending()
        if not q:
            return 0
        self.pending_path.unlink(missing_ok=True)
        done = 0
        for i, item in enumerate(q):
            fn = getattr(self, item["op"], None)
            if fn is None:
                continue
            try:
                fn(**item["payload"])
            except Exception:
                self._write_pending(self._read_pending() + q[i:])
                raise
            if self.pending_path.exists():  # re-queued: still locked
                self._write_pending(self._read_pending() + q[i + 1:])
                break
            done += 1
        return done

    # --- jobs --------------------------------------------------------------

    def _row_data(self, row: TrackerRow | dict[str, Any]) -> dict[str, Any]:
        data = row.model_dump(exclude_none=True) if isinstance(row, TrackerRow) else {k: v for k, v in row.items() if v is not None}
        if "job_id" not in data:
            raise ValueError("job_id required")
        if "folder" not in data and self.jobs_dir is not None:
            data["folder"] = str(self.jobs_dir / data["job_id"])
        return data

    @staticmethod
    def _write_job(ws: Worksheet, hdr: dict[str, int], index: dict[str, int], data: dict[str, Any]) -> str:
        jid = data["job_id"]
        r = index.get(jid)
        created = r is None
        if r is None:
            r = max(index.values(), default=1) + 1
            while ws.cell(row=r, column=hdr["JobID"]).value not in (None, ""):
                r += 1
            index[jid] = r
            _put(ws, r, hdr["Status"], "found")
            _put(ws, r, hdr["DateFound"], _today())
        for header, key, _ in JOB_COLUMNS:
            if key not in data:
                continue
            val = data[key]
            if key == "folder":  # our own job dir path
                cell = _put(ws, r, hdr[header], "open")
                cell.hyperlink = str(val)
                cell.font = Font(color="0563C1", underline="single")
            elif key == "url":  # board data: always text; clickable only for http(s)
                cell = _put(ws, r, hdr[header], str(val))
                cell.hyperlink = str(val) if re.match(r"https?://", str(val), re.I) else None
                if cell.hyperlink is not None:
                    cell.font = Font(color="0563C1", underline="single")
            else:
                _put(ws, r, hdr[header], val)
        return "created" if created else "updated"

    @staticmethod
    def _job_index(ws: Worksheet, hdr: dict[str, int]) -> dict[str, int]:
        col = hdr["JobID"]
        return {
            str(row[0].value): row[0].row
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col)
            if row[0].value not in (None, "")
        }

    def upsert_job(self, row: TrackerRow | dict[str, Any]) -> str:
        data = self._row_data(row)

        def fn(wb: Workbook) -> str:
            ws = wb["Jobs"]
            hdr = _header_index(ws)
            return self._write_job(ws, hdr, self._job_index(ws, hdr), data)

        return self._mutate("upsert_job", {"row": data}, fn) or "queued"

    def upsert_jobs(self, rows: list[TrackerRow | dict[str, Any]]) -> dict[str, int]:
        datas = [self._row_data(r) for r in rows]

        def fn(wb: Workbook) -> dict[str, int]:
            ws = wb["Jobs"]
            hdr = _header_index(ws)
            index = self._job_index(ws, hdr)
            counts = {"created": 0, "updated": 0}
            for d in datas:
                counts[self._write_job(ws, hdr, index, d)] += 1
            return counts

        return self._mutate("upsert_jobs", {"rows": datas}, fn) or {"created": 0, "updated": 0, "queued": len(datas)}

    def set_status(self, job_id: str, status: str, note: str | None = None) -> None:
        if status not in STATUSES:
            raise ValueError(f"invalid status '{status}'; expected one of {STATUSES}")

        def fn(wb: Workbook) -> None:
            ws = wb["Jobs"]
            hdr = _header_index(ws)
            r = _find_row(ws, hdr["JobID"], job_id)
            if r is None:
                r = ws.max_row + 1
                _put(ws, r, hdr["JobID"], job_id)
                _put(ws, r, hdr["DateFound"], _today())
            _put(ws, r, hdr["Status"], status)
            if status == "applied" and not ws.cell(row=r, column=hdr["DateApplied"]).value:
                _put(ws, r, hdr["DateApplied"], _today())
            if note:
                prev = ws.cell(row=r, column=hdr["Notes"]).value
                _put(ws, r, hdr["Notes"], f"{prev}\n{note}" if prev else note)
            self._append_log_ws(wb, job_id, "tracker", f"status -> {status}" + (f": {note}" if note else ""))

        self._mutate("set_status", {"job_id": job_id, "status": status, "note": note}, fn)

    def get_job(self, job_id: str) -> dict[str, Any] | None:
        wb = self._load()
        ws = wb["Jobs"]
        hdr = _header_index(ws)
        r = _find_row(ws, hdr["JobID"], job_id)
        if r is None:
            return None
        return {h: ws.cell(row=r, column=c).value for h, c in hdr.items()}

    def list_jobs(self) -> list[dict[str, Any]]:
        wb = self._load()
        ws = wb["Jobs"]
        hdr = _header_index(ws)
        out = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[hdr["JobID"] - 1] in (None, ""):
                continue
            out.append({h: row[c - 1] for h, c in hdr.items()})
        return out

    def read_overrides(self) -> dict[str, str]:
        return {
            str(j["JobID"]): str(j["Override"]).strip()
            for j in self.list_jobs()
            if j.get("Override") not in (None, "")
        }

    def applied_count(self, company: str | None, days: int = 90) -> int:
        """Rows with a DateApplied within the last `days` days, for one company (normalized name) or,
        with company=None, across all companies (`days=1` = applications made today)."""
        key = normalize_company(company) if company is not None else None
        cutoff = datetime.now() - timedelta(days=days)
        n = 0
        for j in self.list_jobs():
            if key is not None and normalize_company(str(j.get("Company") or "")) != key:
                continue
            d = j.get("DateApplied")
            if d in (None, ""):
                continue
            dt = d if isinstance(d, datetime) else _parse_date(str(d))
            if dt and dt >= cutoff:
                n += 1
        return n

    # --- action items ------------------------------------------------------

    def add_action_item(
        self,
        what: str,
        type: str = "other",
        job_id: str = "",
        company: str = "",
        role: str = "",
        link: str = "",
        priority: str = "M",
        id: str | None = None,
        needs: str = "anytime",
    ) -> str:
        item = ActionItem(
            id=id or uuid.uuid4().hex[:8],
            what=what, type=type, job_id=job_id, company=company, role=role,  # type: ignore[arg-type]
            link=link, priority=priority, needs=needs,  # type: ignore[arg-type]
        )

        def fn(wb: Workbook) -> str:
            ws = wb["Action Items"]
            hdr = _header_index(ws)
            if _find_row(ws, hdr["ID"], item.id) is not None:
                return item.id
            r = ws.max_row + 1
            vals = {"ID": item.id, "Created": item.created[:19].replace("T", " "), "JobID": item.job_id,
                    "Company": item.company, "Role": item.role, "Type": item.type, "What to do": item.what,
                    "Link": item.link, "Priority": item.priority, "Needs": item.needs, "Done": "N", "DoneDate": ""}
            for h, v in vals.items():
                if h in hdr:
                    _put(ws, r, hdr[h], v)
            self._append_log_ws(wb, item.job_id, "action", f"[{item.type}/{item.priority}/{item.needs}] {item.what}")
            return item.id

        payload = item.model_dump(include={"what", "type", "job_id", "company", "role", "link", "priority", "id", "needs"})
        return self._mutate("add_action_item", payload, fn) or item.id

    def list_action_items(self, open_only: bool = True) -> list[dict[str, Any]]:
        wb = self._load()
        ws = wb["Action Items"]
        hdr = _header_index(ws)
        out = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] in (None, ""):
                continue
            d = {h: row[c - 1] for h, c in hdr.items()}
            if open_only and str(d.get("Done") or "N").upper() == "Y":
                continue
            out.append(d)
        return out

    def mark_action_done(self, id: str) -> bool | None:
        """True = marked, False = no such id, None = tracker locked and the op was queued."""
        def fn(wb: Workbook) -> bool:
            ws = wb["Action Items"]
            hdr = _header_index(ws)
            r = _find_row(ws, hdr["ID"], id)
            if r is None:
                return False
            _put(ws, r, hdr["Done"], "Y")
            _put(ws, r, hdr["DoneDate"], _today())
            return True

        return self._mutate("mark_action_done", {"id": id}, fn)

    # --- contacts ----------------------------------------------------------

    def add_contact(
        self,
        company: str,
        name: str,
        job_id: str = "",
        title: str = "",
        linkedin: str = "",
        email: str = "",
        email_confidence: str = "",
        draft_message: str = "",
    ) -> None:
        c = Contact(company=company, name=name, job_id=job_id, title=title, linkedin=linkedin,
                    email=email, email_confidence=email_confidence, draft_message=draft_message)

        def fn(wb: Workbook) -> None:
            ws = wb["Contacts"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[1] == c.company and row[2] == c.name and (row[0] or "") == c.job_id:
                    return
            r = ws.max_row + 1
            vals = [c.job_id, c.company, c.name, c.title, c.linkedin, c.email, c.email_confidence,
                    c.draft_message, "N", "", ""]
            for i, v in enumerate(vals, start=1):
                _put(ws, r, i, v)

        payload = c.model_dump(include={"company", "name", "job_id", "title", "linkedin", "email", "email_confidence", "draft_message"})
        self._mutate("add_contact", payload, fn)

    # --- log / config ------------------------------------------------------

    def _append_log_ws(self, wb: Workbook, job_id: str, component: str, message: str) -> None:
        ws = wb["Log"]
        r = ws.max_row + 1
        for i, v in enumerate([_now(), job_id, component, message], start=1):
            _put(ws, r, i, v)

    def log(self, job_id: str, component: str, message: str) -> None:
        self._mutate("log", {"job_id": job_id, "component": component, "message": message},
                     lambda wb: self._append_log_ws(wb, job_id, component, message))

    def set_config(self, key: str, value: Any) -> None:
        def fn(wb: Workbook) -> None:
            ws = wb["Config"]
            r = _find_row(ws, 1, key)
            if r is None:
                r = ws.max_row + 1
                _put(ws, r, 1, key)
            _put(ws, r, 2, value)

        self._mutate("set_config", {"key": key, "value": value}, fn)

    def get_config(self) -> dict[str, Any]:
        wb = self._load()
        ws = wb["Config"]
        return {str(r[0]): r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}


def _parse_date(s: str) -> datetime | None:
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if not m:
        return None
    try:
        return datetime(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None


def set_status_both(settings: Settings, job_id: str, status: str, note: str) -> None:
    """Set a job's status in data/jobs/<id>/status.json and in the tracker (queued if the tracker is locked)."""
    from careeros.store import Store

    Store(settings).set_status(job_id, status, note)
    Tracker(settings=settings).set_status(job_id, status, note)
